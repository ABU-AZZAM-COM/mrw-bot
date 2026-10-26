"""
اختبار سريع بدون تلغرام: يتحقق من أن كل الوحدات تستورد وتعمل معًا.
(التوكن يُزوَّر عبر متغيرات البيئة خلال الاختبار فقط)
"""
import os, time, tempfile, asyncio, json
import sqlite3
from datetime import datetime, timedelta, timezone
from pathlib import Path
from types import SimpleNamespace

os.environ["BOT_TOKEN"] = "123456:TEST_TOKEN_FOR_IMPORT_CHECK"
os.environ["ADMIN_IDS"] = "111111111"

import config
import guards
import texts
import utils
import database
import keyboards
import exporters
import handlers
import admin
import archive
import health_server
import main as main_mod
import sightengine_moderation
from unittest.mock import AsyncMock, patch

async def _capture_denied_command(*args, **kwargs):
    return None

denied_command_update = SimpleNamespace(
    effective_message=SimpleNamespace(
        text="/admin",
        reply_text=_capture_denied_command,
    ),
    effective_user=SimpleNamespace(id=222222222),
)
try:
    asyncio.run(main_mod.command_access_control(denied_command_update, None))
except main_mod.ApplicationHandlerStop:
    pass
else:
    raise AssertionError("non-admin /admin command was not stopped")
assert asyncio.run(
    main_mod.command_access_control(
        SimpleNamespace(
            effective_message=SimpleNamespace(text="/start", reply_text=_capture_denied_command),
            effective_user=SimpleNamespace(id=222222222),
        ),
        None,
    )
) is None
assert asyncio.run(
    main_mod.command_access_control(
        SimpleNamespace(
            effective_message=SimpleNamespace(text="/unknown", reply_text=_capture_denied_command),
            effective_user=SimpleNamespace(id=111111111),
        ),
        None,
    )
) is None
start_reply = AsyncMock()
start_db = SimpleNamespace(
    record_user=AsyncMock(),
    log_event=AsyncMock(),
    get_state=AsyncMock(return_value=None),
)
start_update = SimpleNamespace(
    effective_user=SimpleNamespace(
        id=222222222, username="start_test", full_name="اختبار البدء"
    ),
    effective_chat=SimpleNamespace(type="private"),
    effective_message=SimpleNamespace(reply_text=start_reply),
)
start_context = SimpleNamespace(
    bot_data={"db": start_db}, user_data={"stale": "draft"}
)
assert asyncio.run(handlers.cmd_start(start_update, start_context)) is (
    handlers.ConversationHandler.END
)
start_reply.assert_awaited_once()
assert start_reply.await_args.args[0] == texts.WELCOME
assert start_context.user_data == {}
print("/start command builds and sends the welcome response")
wizard_states = handlers.build_conversation().states
assert all(
    state in wizard_states
    for state in (
        handlers.ASK_CATEGORY,
        handlers.ASK_PLATFORM,
        handlers.ASK_DESC,
        handlers.ASK_PHOTO,
        handlers.ASK_LINK,
        handlers.ASK_LOCATION,
        handlers.ASK_CONTACT,
        handlers.ASK_FOLLOWUP,
        handlers.ASK_SUPPORT,
        handlers.CONFIRM,
    )
)
assert "ستراجعها الجهة المختصة" in texts.MY_REPORT_OPEN.format(
    seq="TEST",
    time="TEST",
    status="TEST",
    status_detail="ستراجعها الجهة المختصة",
    desc="TEST",
)
prompt_edits = AsyncMock()
prompt_context = SimpleNamespace(
    user_data={},
    bot=SimpleNamespace(edit_message_reply_markup=prompt_edits),
)
handlers._remember_optional_prompt(
    prompt_context, SimpleNamespace(chat_id=12, message_id=34)
)
asyncio.run(handlers._clear_optional_prompt(prompt_context))
assert not prompt_context.user_data
prompt_edits.assert_awaited_once_with(
    chat_id=12, message_id=34, reply_markup=None
)

_at_threshold_scores = sightengine_moderation.parse_scores({
    "status": "success",
    "nudity": {
        "sexual_activity": 0.10,
        "sexual_display": 0.40,
        "erotica": 0.10,
        "suggestive": 0.40,
        "mildly_suggestive": 0.40,
        "none": 0.99,
        "context": {"indoor_other": 0.99},
        "suggestive_classes": {
            "lingerie": 0.40,
            "bikini": 0.40,
            "cleavage_categories": {"very_revealing": 0.40, "none": 0.99},
        },
    },
})
assert not _at_threshold_scores.is_restricted
_suggestive_scores = sightengine_moderation.parse_scores({
    "status": "success",
    "nudity": {
        "sexual_activity": 0.10,
        "sexual_display": 0.10,
        "erotica": 0.10,
        "very_suggestive": 0.4001,
        "suggestive_classes": {
            "lingerie": 0.43,
            "bikini": 0.46,
            "male_underwear": 0.10,
            "minishort": 0.10,
            "suggestive_focus": 0.10,
            "suggestive_pose": 0.10,
        },
    },
})
assert _suggestive_scores.is_restricted
assert _suggestive_scores.max_score == 0.46
_case_image_scores = sightengine_moderation.parse_scores({
    "status": "success",
    "nudity": {
        "sexual_activity": 0.34,
        "sexual_display": 0.50,
        "erotica": 0.05,
    },
})
assert _case_image_scores.is_restricted
assert _case_image_scores.max_score == 0.50
assert sightengine_moderation.parse_scores({
    "status": "success",
    "nudity": {
        "sexual_activity": 0.71,
        "sexual_display": 0.10,
        "erotica": 0.10,
    },
}).is_restricted
assert sightengine_moderation.parse_scores({
    "status": "success",
    "nudity": {
        "sexual_activity": 0.10,
        "sexual_display": 0.10,
        "erotica": 0.71,
    },
}).is_restricted
_other_content_only = sightengine_moderation.parse_scores({
    "status": "success",
    "nudity": {
        "sexual_activity": 0.10,
        "sexual_display": 0.10,
        "erotica": 0.10,
        "context": {"indoor_other": 0.99},
        "none": 0.99,
    },
})
assert not _other_content_only.is_restricted
assert "HTTP 401" in sightengine_moderation._http_status_message(401)
assert "HTTP 403" in sightengine_moderation._http_status_message(403)
assert "رصيد الحساب" in sightengine_moderation._http_status_message(429)
assert "HTTP 503" in sightengine_moderation._http_status_message(503)
try:
    sightengine_moderation.parse_scores({
        "status": "failure",
        "error": {"type": "credentials_error", "code": 1},
    })
except sightengine_moderation.SightengineError as exc:
    assert "بيانات الاعتماد" in str(exc)
else:
    raise AssertionError("Sightengine credential failure should be reported")
try:
    sightengine_moderation.parse_scores({
        "status": "failure",
        "error": {"type": "usage_limit", "code": 2},
    })
except sightengine_moderation.SightengineError as exc:
    assert "حصة أو حد" in str(exc)
else:
    raise AssertionError("Sightengine usage limit should be reported")

_sightengine_request = {}

class _SightengineResponse:
    def raise_for_status(self):
        return None

    def json(self):
        return {
            "status": "success",
            "nudity": {
                "sexual_activity": 0.2,
                "sexual_display": 0.1,
                "erotica": 0.1,
            },
        }

class _SightengineClient:
    def __init__(self, **_kwargs):
        pass

    async def __aenter__(self):
        return self

    async def __aexit__(self, *_args):
        return None

    async def post(self, url, *, data, files):
        _sightengine_request.update(
            url=url,
            data=data,
            filename=files["media"][0],
            content=files["media"][1].read(),
        )
        return _SightengineResponse()

_old_sightengine_user = os.environ.get("SIGHTENGINE_API_USER")
_old_sightengine_secret = os.environ.get("SIGHTENGINE_API_SECRET")
os.environ["SIGHTENGINE_API_USER"] = "test-user"
os.environ["SIGHTENGINE_API_SECRET"] = "test-secret"
with tempfile.TemporaryDirectory(prefix="mrw-sightengine-selftest-") as sightengine_dir:
    sightengine_image = Path(sightengine_dir) / "test.jpg"
    sightengine_image.write_bytes(b"api-test-image")
    with patch.object(sightengine_moderation.httpx, "AsyncClient", _SightengineClient):
        _api_scores = asyncio.run(
            sightengine_moderation.inspect_image(sightengine_image)
        )
assert _api_scores.scores == (
    ("sexual_activity", 0.2),
    ("sexual_display", 0.1),
    ("erotica", 0.1),
)
assert _sightengine_request["url"] == sightengine_moderation.SIGHTENGINE_URL
assert _sightengine_request["data"]["models"] == sightengine_moderation.SIGHTENGINE_MODEL
assert _sightengine_request["content"] == b"api-test-image"
if _old_sightengine_user is None:
    os.environ.pop("SIGHTENGINE_API_USER", None)
else:
    os.environ["SIGHTENGINE_API_USER"] = _old_sightengine_user
if _old_sightengine_secret is None:
    os.environ.pop("SIGHTENGINE_API_SECRET", None)
else:
    os.environ["SIGHTENGINE_API_SECRET"] = _old_sightengine_secret
print("Sightengine API request and nudity-score threshold parsing OK")

_old_sightengine_user = os.environ.get("SIGHTENGINE_API_USER")
_old_sightengine_secret = os.environ.get("SIGHTENGINE_API_SECRET")
os.environ["SIGHTENGINE_API_USER"] = "test-user"
os.environ["SIGHTENGINE_API_SECRET"] = "test-secret"
with tempfile.TemporaryDirectory(prefix="mrw-private-report-selftest-") as image_dir:
    uploaded_image = Path(image_dir) / "report.jpg"
    uploaded_image.write_bytes(b"private-report-image")
    with (
        patch.object(
            handlers.image_moderation,
            "screen_image",
            return_value=handlers.image_moderation.ScreeningResult(False),
        ),
        patch.object(
            sightengine_moderation,
            "inspect_image",
            new=AsyncMock(
                return_value=_case_image_scores
            ),
        ),
    ):
        _private_report_screening = asyncio.run(
            handlers._screen_uploaded_photo(
                media=None,
                photo_path=str(uploaded_image),
                extension=".jpg",
                user_id=654321,
            )
        )
    assert _private_report_screening.flagged
    assert "SIGHTENGINE_SEXUAL_DISPLAY" in _private_report_screening.labels
    assert _private_report_screening.max_score == 0.50
    flagged_result = _private_report_screening

    with (
        patch.object(
            handlers.image_moderation,
            "screen_image",
            return_value=handlers.image_moderation.ScreeningResult(False),
        ),
        patch.object(
            sightengine_moderation,
            "inspect_image",
            new=AsyncMock(return_value=_suggestive_scores),
        ),
    ):
        _suggestive_report_screening = asyncio.run(
            handlers._screen_uploaded_photo(
                media=None,
                photo_path=str(uploaded_image),
                extension=".jpg",
                user_id=654321,
            )
        )
    assert _suggestive_report_screening.flagged
    assert "SIGHTENGINE_SUGGESTIVE_CLASSES_LINGERIE" in (
        _suggestive_report_screening.labels
    )
    assert "SIGHTENGINE_SUGGESTIVE_CLASSES_BIKINI" in (
        _suggestive_report_screening.labels
    )

    with (
        patch.object(
            handlers.image_moderation,
            "screen_image",
            return_value=handlers.image_moderation.ScreeningResult(False),
        ),
        patch.object(
            sightengine_moderation,
            "inspect_image",
            new=AsyncMock(
                return_value=sightengine_moderation.NudityScores(
                    scores=(
                        ("sexual_activity", 0.10),
                        ("sexual_display", 0.40),
                        ("erotica", 0.10),
                    )
                )
            ),
        ),
    ):
        _threshold_report_screening = asyncio.run(
            handlers._screen_uploaded_photo(
                media=None,
                photo_path=str(uploaded_image),
                extension=".jpg",
                user_id=654321,
            )
        )
    assert not _threshold_report_screening.flagged

    with (
        patch.object(
            handlers.image_moderation,
            "screen_image",
            return_value=handlers.image_moderation.ScreeningResult(False),
        ),
        patch.object(
            sightengine_moderation,
            "inspect_image",
            new=AsyncMock(
                return_value=_case_image_scores
            ),
        ),
    ):
        _over_threshold_report_screening = asyncio.run(
            handlers._screen_uploaded_photo(
                media=None,
                photo_path=str(uploaded_image),
                extension=".jpg",
                user_id=654321,
            )
        )
    assert _over_threshold_report_screening.flagged
    assert "SIGHTENGINE_SEXUAL_DISPLAY" in _over_threshold_report_screening.labels

    class _PrivateTelegramFile:
        async def download_to_drive(self, custom_path):
            Path(custom_path).write_bytes(b"private-temporary-image")

    class _PrivateMedia:
        async def get_file(self):
            return _PrivateTelegramFile()

    _temporary_scan_paths = []

    async def _inspect_temporary_image(path):
        _temporary_scan_paths.append(Path(path))
        assert Path(path).read_bytes() == b"private-temporary-image"
        return sightengine_moderation.NudityScores(
            scores=(
                ("sexual_activity", 0.1),
                ("sexual_display", 0.1),
                ("erotica", 0.1),
            )
        )

    with (
        patch.object(
            handlers.image_moderation,
            "screen_image",
            return_value=handlers.image_moderation.ScreeningResult(False),
        ),
        patch.object(
            sightengine_moderation,
            "inspect_image",
            new=_inspect_temporary_image,
        ),
    ):
        asyncio.run(
            handlers._screen_uploaded_photo(
                media=_PrivateMedia(),
                photo_path="",
                extension=".jpg",
                user_id=654321,
            )
        )
    assert _temporary_scan_paths and all(
        not path.exists() for path in _temporary_scan_paths
    )
if _old_sightengine_user is None:
    os.environ.pop("SIGHTENGINE_API_USER", None)
else:
    os.environ["SIGHTENGINE_API_USER"] = _old_sightengine_user
if _old_sightengine_secret is None:
    os.environ.pop("SIGHTENGINE_API_SECRET", None)
else:
    os.environ["SIGHTENGINE_API_SECRET"] = _old_sightengine_secret
print("private complaint image: Sightengine >40% nudity score freezes")

_db_test_dir = tempfile.TemporaryDirectory(prefix="mrw-selftest-")
DB_PATH = Path(_db_test_dir.name) / "test_reports.db"
DB = database.Database(DB_PATH)
assert DB.backend == "sqlite"
assert DB.health_check()
assert DB._conn.execute("SELECT 1 AS value").fetchone()[0] == 1

notifications_db = database.Database(
    Path(_db_test_dir.name) / "notifications.db"
)
notification_subscriber = notifications_db.subscribe_notifications()
notification_report = asyncio.run(notifications_db.add_report(
    user_id=7001, username="notify_test", full_name="مستخدم إشعار",
    photo_path="", photo_file_id=None, post_link="https://example.com/notify",
    description="بلاغ لاختبار الإشعار",
))
report_event = notification_subscriber.get_nowait()
assert report_event["type"] == "report_new"
assert report_event["payload"]["report_id"] == notification_report["id"]
asyncio.run(notifications_db.add_user_message(7001, "in", "رسالة واردة"))
message_event = notification_subscriber.get_nowait()
assert message_event["type"] == "user_message"
asyncio.run(notifications_db.save_block(
    7001, 1, "9999-12-31T00:00:00", "اختبار إشعار الحظر"
))
block_event = notification_subscriber.get_nowait()
assert block_event["type"] == "user_blocked"
asyncio.run(notifications_db.update_report_workflow(
    notification_report["id"], "urgent", None, None
))
urgent_event = notification_subscriber.get_nowait()
assert urgent_event["type"] == "report_urgent"
asyncio.run(notifications_db.set_state("emergency", "1"))
emergency_event = notification_subscriber.get_nowait()
assert emergency_event["type"] == "emergency"
items, unread_count = asyncio.run(notifications_db.list_notifications(limit=20))
assert unread_count == 5 and len(items) == 5
assert asyncio.run(
    notifications_db.mark_notification_read(items[0]["id"])
)
assert asyncio.run(notifications_db.mark_all_notifications_read()) == 4
notifications_db.unsubscribe_notifications(notification_subscriber)
notifications_db.close()
print("notifications: persisted events, cross-thread subscription, and read state")

import turso_serverless

mock_turso_connection = sqlite3.connect(":memory:", check_same_thread=False)
with (
    patch.dict(
        os.environ,
        {
            "TURSO_DATABASE_URL": "libsql://unit-test.example",
            "TURSO_AUTH_TOKEN": "unit-test-token",
        },
    ),
    patch.object(turso_serverless, "connect", return_value=mock_turso_connection) as connect_turso,
):
    turso_db = database.Database(Path("unused-turso-local.db"), backend="turso")
    assert turso_db.backend == "turso"
    assert turso_db.health_check()
    assert turso_db._conn.execute("SELECT 1 AS value").fetchone()[0] == 1
    assert connect_turso.call_args.args == ("libsql://unit-test.example",)
    assert connect_turso.call_args.kwargs == {"auth_token": "unit-test-token"}
    assert asyncio.run(turso_db.total()) == 0
    asyncio.run(turso_db.add_report(
        user_id=111, username="snapshot_test", full_name="Snapshot test",
        photo_path="", photo_file_id=None, post_link="https://example.com/report",
        description="snapshot row",
    ))
    turso_snapshot = Path(_db_test_dir.name) / "turso-snapshot.db"
    turso_db.export_sqlite_snapshot(turso_snapshot)
    snapshot_conn = sqlite3.connect(turso_snapshot)
    try:
        assert snapshot_conn.execute(
            "SELECT description FROM reports WHERE user_id = 111"
        ).fetchone() == ("snapshot row",)
    finally:
        snapshot_conn.close()
    turso_db.close()

with patch.dict(os.environ, {"RENDER_SERVICE_NAME": "mrw-test", "DB_BACKEND": "sqlite"}):
    try:
        database.Database(Path("must-not-open-sqlite.db"), backend="sqlite")
    except RuntimeError as error:
        assert "Render production requires DB_BACKEND=turso" in str(error)
    else:
        raise AssertionError("Render must reject SQLite")

with patch.dict(os.environ, {"TURSO_DATABASE_URL": "", "TURSO_AUTH_TOKEN": ""}):
    missing_turso_path = Path("must-not-create-local-fallback.db")
    try:
        database.Database(missing_turso_path, backend="turso")
    except RuntimeError as error:
        assert "TURSO_DATABASE_URL" in str(error)
        assert "TURSO_AUTH_TOKEN" in str(error)
    else:
        raise AssertionError("Turso must fail when credentials are missing")
    assert not missing_turso_path.exists()

support_user = SimpleNamespace(id=81818181, username="support_test", full_name="مستخدم الدعم")
support_reply = AsyncMock()
support_send = AsyncMock()
support_db_dir = tempfile.TemporaryDirectory(prefix="mrw-support-selftest-")
support_db = database.Database(Path(support_db_dir.name) / "support.db")
support_update = SimpleNamespace(
    effective_user=support_user,
    effective_message=SimpleNamespace(text="رسالة اختبار", reply_text=support_reply),
)
support_context = SimpleNamespace(
    user_data={"support_pending": True},
    bot_data={"db": support_db},
    bot=SimpleNamespace(send_message=support_send),
)
support_state = asyncio.run(
    handlers.receive_support_message(support_update, support_context)
)
assert support_state == handlers.ConversationHandler.END
assert not support_context.user_data
assert asyncio.run(support_db.user_conversation_open(support_user.id))
support_messages = asyncio.run(support_db.list_user_messages(support_user.id))
assert len(support_messages) == 1
assert support_messages[0]["text"] == "رسالة اختبار"
support_send.assert_awaited_once()
support_reply.assert_awaited_once()
support_db.close()
support_db_dir.cleanup()

freeze_alerts = []
freeze_notices = []

async def _capture_freeze_alert(**kwargs):
    freeze_alerts.append(kwargs)

async def _capture_freeze_notice(text, **kwargs):
    freeze_notices.append((text, kwargs))

freeze_user = SimpleNamespace(id=90909090, full_name="مستخدم اختبار")
freeze_update = SimpleNamespace(
    effective_user=freeze_user,
    effective_message=SimpleNamespace(reply_text=_capture_freeze_notice),
)
freeze_context = SimpleNamespace(
    bot_data={"db": DB},
    bot=SimpleNamespace(send_message=_capture_freeze_alert),
    user_data={"app_draft": {"description": "should be discarded"}},
)
handlers._pending_app_users.add(freeze_user.id)
asyncio.run(
    handlers._freeze_for_sensitive_image(
        freeze_update,
        freeze_context,
        handlers.image_moderation.ScreeningResult(
            True,
            review_required=True,
            labels=("FEMALE_GENITALIA_EXPOSED",),
            max_score=0.97,
        ),
        photo_path="media/sensitive-review.jpg",
        photo_file_id="test-sensitive-file-id",
    )
)
automatic_block = asyncio.run(DB.get_block(freeze_user.id))
assert automatic_block["frozen_until"] == "9999-12-31T00:00:00"
assert automatic_block["reason"].startswith("كشف آلي")
assert asyncio.run(DB.get_block_evidence(freeze_user.id)) == {
    "photo_path": "media/sensitive-review.jpg",
    "photo_file_id": "test-sensitive-file-id",
}
assert not freeze_context.user_data and freeze_user.id not in handlers._pending_app_users
assert len(freeze_alerts) == len(config.ADMIN_IDS) and len(freeze_notices) == 1
assert all("photo" not in alert for alert in freeze_alerts)
assert any(
    row["action"] == "automated_sensitive_image_freeze"
    and row["target"] == str(freeze_user.id)
    for row in asyncio.run(DB.recent_audit(20))
)
blocked_rows, _ = asyncio.run(DB.list_blocked_users())
assert any(row["user_id"] == freeze_user.id for row in blocked_rows)
blocked_verdict = asyncio.run(guards.check_rate_limits(DB, freeze_user.id))
assert blocked_verdict.blocked, blocked_verdict
asyncio.run(DB.clear_block(freeze_user.id))
assert asyncio.run(DB.get_block_evidence(freeze_user.id)) is None
assert asyncio.run(guards.check_rate_limits(DB, freeze_user.id)).allowed
print("automatic explicit-image freeze persists, audits, alerts without forwarding image, and admin unblock restores access")

original_save_photo = handlers._save_photo

async def _fake_flagged_photo(_update, _context):
    return "", "test-file-id", "test-file-id", flagged_result

async def _capture_progress_edit(*args, **kwargs):
    return None

async def _capture_progress_reply(*args, **kwargs):
    return SimpleNamespace(edit_text=_capture_progress_edit)

handlers._save_photo = _fake_flagged_photo
try:
    conversation_user = SimpleNamespace(id=90909091, full_name="مستخدم محادثة")
    conversation_context = SimpleNamespace(
        bot_data={"db": DB},
        bot=SimpleNamespace(send_message=_capture_freeze_alert),
        user_data={"description": "must be discarded"},
    )
    conversation_update = SimpleNamespace(
        message=SimpleNamespace(reply_text=_capture_progress_reply),
        effective_user=conversation_user,
        effective_message=SimpleNamespace(reply_text=_capture_freeze_notice),
    )
    conversation_state = asyncio.run(
        handlers.receive_photo(conversation_update, conversation_context)
    )
    assert conversation_state == handlers.ConversationHandler.END

    app_user = SimpleNamespace(id=90909092, full_name="مستخدم التطبيق")
    app_context = SimpleNamespace(
        bot_data={"db": DB},
        bot=SimpleNamespace(send_message=_capture_freeze_alert),
        user_data={"app_draft": {"post_link": "https://example.test/post"}},
    )
    handlers._pending_app_users.add(app_user.id)
    app_reply_calls = [0]

    async def _capture_app_reply(*args, **kwargs):
        app_reply_calls[0] += 1
        if app_reply_calls[0] == 1:
            return SimpleNamespace(edit_text=_capture_progress_edit)
        return await _capture_freeze_notice(*args, **kwargs)

    app_update = SimpleNamespace(
        message=SimpleNamespace(),
        effective_user=app_user,
        effective_message=SimpleNamespace(reply_text=_capture_app_reply),
    )
    assert asyncio.run(handlers.app_photo_received(app_update, app_context)) is None
    assert not asyncio.run(DB.list_user_reports(conversation_user.id))[0]
    assert not asyncio.run(DB.list_user_reports(app_user.id))[0]
    assert not app_context.user_data and app_user.id not in handlers._pending_app_users
finally:
    handlers._save_photo = original_save_photo
print("conversation and Mini App photo flows freeze and stop without saving a report")

fake_row = dict(
    user_id=999,
    username="tester",
    full_name="مجرم؟ لا، مبلّغ تجريبي",
    photo_path=r"media/fake.jpg",
    photo_file_id="fid",
    post_link="https://t.me/somechannel/123",
    description="منشور يحتوي محتوى مخالف =:= خطير",
)

r1 = asyncio.run(DB.add_report(**fake_row))
r2 = asyncio.run(DB.add_report(**fake_row))
assert r1["status"] == "جديد" and r1["seq"].startswith("MRV-"), r1
print("add_report OK:", r1["seq"], "|", r2["seq"])

review_photo_path = Path(_db_test_dir.name) / "review-image.jpg"
review_photo_path.write_bytes(b"test-image")
review_report = {
    **r2,
    "photo_path": str(review_photo_path),
    "photo_file_id": "review-file-id",
    "admin_note": "مؤشر تعرٍّ يحتاج مراجعة",
    "needs_review": 1,
}
review_notifications = []

async def _capture_review_photo(**kwargs):
    review_notifications.append(("photo", kwargs))

async def _capture_review_message(**kwargs):
    review_notifications.append(("message", kwargs))

review_context = SimpleNamespace(
    bot=SimpleNamespace(
        send_photo=_capture_review_photo,
        send_message=_capture_review_message,
    )
)
asyncio.run(handlers._notify_admins(review_context, review_report))
assert len(review_notifications) == len(config.ADMIN_IDS)
assert all(kind == "photo" for kind, _ in review_notifications)
assert all(
    "مؤشر تعرٍّ يحتاج مراجعة" in kwargs["caption"]
    for _, kwargs in review_notifications
)
print("review-required reports notify admins with the photo and moderation note")

rows, total = asyncio.run(DB.list_reports(None, 0, 5))
assert total == 2 and rows[0]["seq"] == r2["seq"], (total, rows)
print("list_reports OK, total =", total)

ok = asyncio.run(DB.update_status(r2["id"], "قيد المراجعة", admin_id=111111111))
assert ok, "update_status failed"
r2b = asyncio.run(DB.get_by_id(r2["id"]))
assert r2b["status"] == "قيد المراجعة" and r2b["reviewed_by"] == 111111111
print("update_status OK ->", r2b["status"])

workflow_due = (datetime.now(timezone.utc) - timedelta(hours=1)).isoformat(
    timespec="seconds"
)
assert asyncio.run(
    DB.update_report_workflow(r1["id"], "urgent", 111111111, workflow_due)
)
workflow_report = asyncio.run(DB.get_by_id(r1["id"]))
assert workflow_report["priority"] == "urgent"
assert workflow_report["assigned_to"] == 111111111
assert workflow_report["due_at"] == workflow_due
assert asyncio.run(DB.workflow_counts()) == {
    "urgent": 1, "overdue": 1, "unassigned": 1,
}
asyncio.run(DB.set_runtime_settings({"MAX_DESC_LEN": 1800}))
assert asyncio.run(DB.get_runtime_settings()) == {"MAX_DESC_LEN": 1800}
print("report workflow and persisted runtime settings OK")

legacy_tmp = tempfile.TemporaryDirectory(
    prefix="mrw-legacy-selftest-", ignore_cleanup_errors=True
)
legacy_path = Path(legacy_tmp.name) / "legacy.db"
with sqlite3.connect(legacy_path) as legacy_conn:
    legacy_conn.execute(
        "CREATE TABLE reports (id INTEGER PRIMARY KEY, seq TEXT, user_id INTEGER, "
        "username TEXT, full_name TEXT, photo_path TEXT NOT NULL, photo_file_id TEXT, "
        "post_link TEXT NOT NULL, description TEXT NOT NULL, status TEXT NOT NULL, "
        "created_at TEXT NOT NULL, reviewed_by INTEGER, reviewed_at TEXT, admin_note TEXT, "
        "violation_type TEXT, place TEXT, latitude REAL, longitude REAL, source TEXT, "
        "platform TEXT, needs_review INTEGER, published INTEGER)"
    )
    legacy_conn.execute(
        "INSERT INTO reports (id, seq, user_id, photo_path, post_link, description, "
        "status, created_at) VALUES (1, 'LEGACY-1', 123, '', '', 'قديم', 'جديد', ?)",
        (utils.utc_now_iso(),),
    )
legacy_db = database.Database(legacy_path)
legacy_report = asyncio.run(legacy_db.get_by_id(1))
assert legacy_report["seq"] == "LEGACY-1"
assert legacy_report["priority"] == "normal"
assert legacy_report["assigned_to"] is None
legacy_db.close()
legacy_tmp.cleanup()
print("legacy database migration preserves existing reports")

counts = asyncio.run(DB.count_by_status())
tot = asyncio.run(DB.total())
assert tot == 2 and counts.get("جديد") == 1 and counts.get("قيد المراجعة") == 1, (counts, tot)
print("stats OK:", counts, "total =", tot)

assert asyncio.run(DB.total_users()) == 1
asyncio.run(DB.log_event(999, "start"))
assert asyncio.run(DB.total_users()) == 1
asyncio.run(DB.log_event(1000, "start"))
assert asyncio.run(DB.total_users()) == 2
stats_replies = []

async def _capture_stats(text, **kwargs):
    stats_replies.append(text)

stats_update = SimpleNamespace(
    effective_user=SimpleNamespace(id=111111111),
    effective_chat=SimpleNamespace(type="private"),
    effective_message=SimpleNamespace(reply_text=_capture_stats),
)
asyncio.run(
    admin.cmd_stats(stats_update, SimpleNamespace(bot_data={"db": DB}))
)
assert len(stats_replies) == 1 and "المستخدمون الفريدون: <b>2</b>" in stats_replies[0]
stats_replies.clear()
non_admin_update = SimpleNamespace(
    effective_user=SimpleNamespace(id=222222222),
    effective_chat=SimpleNamespace(type="private"),
    effective_message=SimpleNamespace(reply_text=_capture_stats),
)
asyncio.run(
    admin.cmd_stats(non_admin_update, SimpleNamespace(bot_data={"db": DB}))
)
assert len(stats_replies) == 1 and "وصول مرفوض" in stats_replies[0]
stats_replies.clear()
group_admin_update = SimpleNamespace(
    effective_user=SimpleNamespace(id=111111111),
    effective_chat=SimpleNamespace(type="group"),
    effective_message=SimpleNamespace(reply_text=_capture_stats),
)
asyncio.run(
    admin.cmd_stats(group_admin_update, SimpleNamespace(bot_data={"db": DB}))
)
assert len(stats_replies) == 1 and "المحادثة الخاصة" in stats_replies[0]
print("/stats counts unique report and start users, private admin only")

n24 = asyncio.run(
    DB.count_since((datetime.now(timezone.utc) - timedelta(hours=24)).isoformat())
)
day = asyncio.run(DB.count_since(utils.local_day_start_iso(), user_id=999))
assert n24 == 2 and day == 2, (n24, day)
print("count_since OK (24h =", n24, "| today per-user =", day, ")")

import space_app

_space_app = space_app.build_app()
_vtypes_route = next(
    route for route in _space_app.routes
    if getattr(route, "path", None) == "/api/public/vtypes"
)
_prior_bot_app = main_mod.app
try:
    main_mod.app = SimpleNamespace(bot_data={"db": DB})
    _vtypes_response = asyncio.run(_vtypes_route.endpoint())
    _vtypes_payload = json.loads(_vtypes_response.body)
    assert _vtypes_response.status_code == 200
    assert _vtypes_payload["items"] == asyncio.run(DB.effective_vtypes())
    assert _vtypes_response.headers["cache-control"] == "no-store"
finally:
    main_mod.app = _prior_bot_app
print("Mini App public categories endpoint uses live database values")

exports = Path(tempfile.mkdtemp())
xlsx = exporters.export_excel(asyncio.run(DB.all_reports()), out_dir=exports)
txt = exporters.export_text(asyncio.run(DB.all_reports()), out_dir=exports)
assert xlsx.exists() and txt.exists() and xlsx.stat().st_size > 4000 and txt.stat().st_size > 500
from openpyxl import load_workbook

excel_row = dict(asyncio.run(DB.all_reports())[0])
excel_row["description"] = "قبل\x00وسط\x1fبعد"
sanitized_xlsx = exporters.export_excel([excel_row], out_dir=exports)
workbook = load_workbook(sanitized_xlsx, read_only=True)
try:
    description_cell = next(workbook["البلاغات"].iter_rows(
        min_row=2, max_row=2, min_col=8, max_col=8, values_only=True
    ))[0]
    assert description_cell == "قبلوسطبعد"
finally:
    workbook.close()
print("exporters OK:", xlsx.name, "|", txt.name)
print("exporters: Excel strips XML-illegal control characters from report data")

# فحص الترحيب/العرض (نصوص جاهزة) — الترحيب يؤدي إلى القائمة ذات الأزرار
w = texts.WELCOME
v = texts.badge("قيد المراجعة")
assert (
    "منظومة البلاغات والشكاوى" in w
    and "اختر الخدمة التي تحتاجها" in w
    and v == "🟠 قيد المراجعة"
)

# القائمة الرئيسية: 4 إجراءات بلا زر فتح التطبيق + زر الرجوع يعيد إظهارها
mm = keyboards.main_menu_kb()
actions = [b.callback_data for row in mm.inline_keyboard for b in row]
assert actions == ["usr:report", "usr:status", "usr:about", "usr:contact"], actions
back = keyboards.back_to_menu_kb()
assert back.inline_keyboard[0][0].callback_data == "usr:menu"
print("texts OK")

# ------------------------------------------------------------------
# فحص طبقة الأرشفة (بدون اتصال بتليغرام)
# ------------------------------------------------------------------
cap = texts.ARCHIVE_REPORT.format(
    seq="MRV-2026-0001",
    name="مبلّغ تجريبي",
    username="@tester",
    time="2026-10-22 09:00",
    anchor="<a href=\"https://t.me/x/1\">المنشور</a>",
    desc="وصف المخالفة",
)
assert "MRV-2026-0001" in cap and "وصف المخالفة" in cap, cap
assert "<b>" in cap and "{seq}" not in cap, cap
print("ARCHIVE_REPORT OK")

cid = texts.CHAT_ID.format(chat_id=-100123, title="قناة", chat_type="channel", user_id=5)
assert "-100123" in cid, cid
print("CHAT_ID OK")

# النسخ الاحتياطي locally (بلا قناة) يجب أن يعطي False بلا أخطاء
assert asyncio.run(archive.backup_database(None, DB_PATH)) is False
assert asyncio.run(archive.post_report(None, r1)) is False
print("archive: بلا ARCHIVE_CHAT_ID لا يفعل شيئًا (سلوك صحيح)")

# نسخة متسقة من القاعدة + فحص وجود صفوف
snap = Path(tempfile.mkdtemp()) / "snap.db"
assert archive._snapshot(DB_PATH, snap) and snap.exists()
assert archive._has_rows(snap) is True
print("archive: النسخ المتسق وفحص الصفوف OK")

# ------------------------------------------------------------------
# فحص طبقة النسخ إلى Hub (منطق خالص — بلا أي اتصال)
# ------------------------------------------------------------------
import hub_backup

import config as _cfg

_tok, _repo = _cfg.HF_BACKUP_TOKEN, _cfg.HF_BACKUP_REPO
try:
    # غير مهيأ → كل المسارات ترفض فورًا بلا استثناء
    _cfg.HF_BACKUP_TOKEN = ""
    _cfg.HF_BACKUP_REPO = ""
    assert hub_backup.is_configured() is False
    assert hub_backup.restore_latest_if_needed(DB_PATH) is False
    assert hub_backup.upload_backup(DB_PATH) is False

    # مهيأ + القاعدة المحلية فيها بلاغات → لِماذا؟ الحماية الصارمة:
    # لا استعادة فوق قاعدة حية (يجب ألا يُلمس الملف أبدًا) — وبلا اتصال
    _cfg.HF_BACKUP_TOKEN = "fake-token-for-test"
    _cfg.HF_BACKUP_REPO = "fake/repo"
    assert hub_backup.is_configured() is True
    assert archive._has_rows(DB_PATH) is True
    assert hub_backup.restore_latest_if_needed(DB_PATH) is False
finally:
    _cfg.HF_BACKUP_TOKEN, _cfg.HF_BACKUP_REPO = _tok, _repo
print("hub_backup: حرّاس (غير مهيأ / لا استبدال فوق بيانات) OK")

# ------------------------------------------------------------------
# فحص طبقة واجهة الإدارة (منطق خالص — بلا FastAPI وبلا شبكة)
# ------------------------------------------------------------------
import admin_api
import os as _os

# الحماية بالتوكن: فراغ = مرفوض، عدم مطابقة = مرفوض، مطابقة = مقبول
_os.environ["ADMIN_API_TOKEN"] = "test-token-xyz"
assert admin_api.expected_token() == "test-token-xyz"
assert admin_api.token_ok("", "test-token-xyz") is False
assert admin_api.token_ok("wrong", "test-token-xyz") is False
assert admin_api.token_ok("  test-token-xyz  ", "test-token-xyz") is True
_os.environ.pop("ADMIN_API_TOKEN")
assert admin_api.expected_token() == ""
print("admin_api: حماية التوكن OK")

_os.environ["ADMIN_API_TOKENS_JSON"] = json.dumps([
    {"token": "reviewer-test-token", "role": "reviewer", "admin_id": 111111111},
    {"token": "operator-test-token", "role": "operator", "admin_id": 222222222},
])

class _RoleRequest:
    headers = {"x-admin-token": "reviewer-test-token"}
    query_params = {}
    method = "POST"
    state = SimpleNamespace()

    def __init__(self, path):
        self.url = SimpleNamespace(path=path)

class _FileAuthRequest:
    method = "GET"
    state = SimpleNamespace()

    def __init__(self, path, headers, query_params=None):
        self.url = SimpleNamespace(path=path)
        self.headers = headers
        self.query_params = query_params or {}

review_req = _RoleRequest("/api/admin/reports/1/status")
assert admin_api._auth(review_req) is None
assert review_req.state.admin_role == "reviewer"
assert review_req.state.admin_id == 111111111
denied_req = _RoleRequest("/api/admin/settings")
denied = admin_api._auth(denied_req)
assert denied is not None and denied.status_code == 403
for file_path in (
    "/api/admin/photo/1",
    "/api/admin/export/xlsx",
    "/api/admin/backup/download",
):
    query_only = _FileAuthRequest(
        file_path, {}, {"token": "reviewer-test-token"}
    )
    denied = admin_api._auth(query_only)
    assert denied is not None and denied.status_code == 401
    legacy_header = _FileAuthRequest(
        file_path, {"x-admin-token": "reviewer-test-token"}
    )
    denied = admin_api._auth(legacy_header)
    assert denied is not None and denied.status_code == 401
    bearer_header = _FileAuthRequest(
        file_path, {"authorization": "Bearer operator-test-token"}
    )
    assert admin_api._auth(bearer_header) is None
_os.environ.pop("ADMIN_API_TOKENS_JSON")
print("admin_api: server-side role authorization OK")
print("admin_api: file endpoints require Bearer auth; query tokens rejected")

# أزرار رفض البلاغ من إشعار البوت والتفاصيل تطلب السبب.
notify_callbacks = [
    button.callback_data
    for row in keyboards.admin_notify_buttons(123).inline_keyboard
    for button in row
]
detail_callbacks = [
    button.callback_data
    for row in keyboards.admin_view_buttons(123).inline_keyboard
    for button in row
]
assert "adm:rj:123" in notify_callbacks
assert "adm:rj:123" in detail_callbacks
assert "adm:s:123:rejected" not in detail_callbacks
category_callbacks = [
    button.callback_data
    for row in keyboards.report_category_kb().inline_keyboard
    for button in row
]
platform_callbacks = [
    button.callback_data
    for row in keyboards.social_platform_kb().inline_keyboard
    for button in row
]
followup_callbacks = [
    button.callback_data
    for row in keyboards.report_open_kb(42, True).inline_keyboard
    for button in row
]
assert all(
    f"usr:category:{key}" in category_callbacks
    for key in ("conduct", "administrative", "social", "field")
)
assert "usr:platform:whatsapp" in platform_callbacks
assert "usr:followup:42" in followup_callbacks
print("keyboards: الرفض من أزرار البوت يطلب السبب دائمًا")

# أدوات اللوحة على قاعدة اختبار فعلية (محلية مؤقتة)
pbase_tmp = tempfile.TemporaryDirectory(prefix="mrw-panel-selftest-")
pbase = Path(pbase_tmp.name) / "panel.db"
pdb = database.Database(pbase)
pid = asyncio.run(pdb.add_report(
    user_id=777, username="panel_user", full_name="مستخدم لوحة",
    photo_path="media/x.jpg", photo_file_id="fid777",
    contact_phone="+15551234567",
    post_link="https://t.me/x/9",
    description="بلاغ لتجربة واجهة الإدارة الكاملة هنا بالتفصيل الشديد",
    violation_type="مخالفة زي",
))
rows, total = asyncio.run(pdb.list_reports_filtered(user_id=777))
assert total == 1 and rows[0]["id"] == pid["id"]
rows, total = asyncio.run(pdb.list_reports_filtered(q="واجهة الإدارة"))
assert total >= 1 and rows[0]["user_id"] == 777
rows, total = asyncio.run(pdb.list_reports_filtered(vtype="مخالفة زي"))
assert total >= 1
assert asyncio.run(pdb.set_violation_type(pid["id"], "أخرى")) is True
assert asyncio.run(pdb.set_admin_note(pid["id"], "ملاحظة من اللوحة")) is True
check = asyncio.run(pdb.get_by_id(pid["id"]))
assert check["violation_type"] == "أخرى" and check["admin_note"] == "ملاحظة من اللوحة"

# قائمة المستخدمين المجملة + بطاقة مستخدم + المستخدم الوهمي = None
users_rows, users_total = asyncio.run(pdb.list_users())
assert users_total == 1 and users_rows[0]["user_id"] == 777
profile = asyncio.run(pdb.user_profile(777))
assert profile["user_id"] == 777 and profile["total_reports"] == 1
assert profile["contact_phone"] == "+15551234567"
assert profile["contact_phone_seq"] == pid["seq"]
assert asyncio.run(pdb.user_profile(42424242)) is None
asyncio.run(pdb.record_user(778, "subscriber_only", "مشترك بلا بلاغ"))
assert asyncio.run(pdb.total_users()) == 2
subscriber_ids, overflow = asyncio.run(pdb.all_user_ids())
assert not overflow and set(subscriber_ids) == {777, 778}

# الحجب: حفظ ثم ظهور في القائمة ثم الرفع ثم الاختفاء
asyncio.run(pdb.save_block(
    777,
    0,
    "9999-12-31T00:00:00",
    "اختبار لوحة",
    evidence_photo_path="media/sensitive-evidence.jpg",
    evidence_photo_file_id="telegram-evidence-file",
))
users_rows, _ = asyncio.run(pdb.list_users())
blocked_777 = next(row for row in users_rows if row["user_id"] == 777)
assert blocked_777["frozen_until"] == "9999-12-31T00:00:00"
blocked_rows, blocked_total = asyncio.run(pdb.list_blocked_users())
assert blocked_total == 1
assert blocked_rows[0]["user_id"] == 777
assert blocked_rows[0]["reason"] == "اختبار لوحة"
assert blocked_rows[0]["has_evidence_photo"] == 1
assert asyncio.run(pdb.get_block_evidence(777)) == {
    "photo_path": "media/sensitive-evidence.jpg",
    "photo_file_id": "telegram-evidence-file",
}
asyncio.run(pdb.save_block(42424242, 0, "9999-12-31T00:00:00", "بلا بلاغات"))
blocked_rows, blocked_total = asyncio.run(pdb.list_blocked_users())
assert blocked_total == 2
assert {row["user_id"] for row in blocked_rows} == {777, 42424242}
first_page, page_total = asyncio.run(pdb.list_blocked_users(page=0, size=1))
second_page, _ = asyncio.run(pdb.list_blocked_users(page=1, size=1))
assert page_total == 2 and len(first_page) == len(second_page) == 1
assert first_page[0]["user_id"] != second_page[0]["user_id"]
asyncio.run(pdb.save_block(52525252, 0, "2000-01-01T00:00:00", "منتهي"))
active_rows, active_total = asyncio.run(pdb.list_blocked_users())
assert active_total == 2
assert 52525252 not in {row["user_id"] for row in active_rows}
asyncio.run(pdb.log_event(777, "start"))
reset_test_since = "2000-01-01T00:00:00+00:00"
assert asyncio.run(
    pdb.count_events(reset_test_since, user_id=777, kind="start")
) == 1
asyncio.run(pdb.clear_block(42424242))
asyncio.run(pdb.clear_block(777))
asyncio.run(pdb.clear_block(52525252))
assert asyncio.run(pdb.get_block(777)).get("frozen_until") in (None, "")
assert asyncio.run(pdb.get_block_evidence(777)) is None
assert asyncio.run(
    pdb.count_events(reset_test_since, user_id=777, kind="start")
) == 0
assert asyncio.run(pdb.count_events(reset_test_since, kind="start")) == 1
assert asyncio.run(pdb.count_since(utils.local_day_start_iso(), user_id=777)) == 0
assert asyncio.run(pdb.count_since(utils.local_day_start_iso())) >= 1
assert asyncio.run(pdb.list_user_reports(777))[0]

# محادثة المستخدم تحفظ الرسائل في القاعدة وتستعيدها بعد إعادة فتحها.
asyncio.run(pdb.open_user_conversation(777))
asyncio.run(pdb.add_user_message(777, "out", "رسالة من الإدارة", 111))
asyncio.run(pdb.add_user_message(777, "in", "رد المستخدم"))
conversation_db = database.Database(pbase)
conversation = asyncio.run(conversation_db.list_user_messages(777))
assert [message["direction"] for message in conversation] == ["out", "in"]
assert conversation[-1]["text"] == "رد المستخدم"
conversation_db.close()

retention_now = datetime.now(timezone.utc)
old_timestamp = (retention_now - timedelta(days=400)).isoformat(timespec="seconds")
recent_timestamp = (retention_now - timedelta(days=20)).isoformat(timespec="seconds")
with pdb._lock:
    pdb._conn.execute(
        "INSERT INTO audit_log (ts, actor_id, action, target, details) "
        "VALUES (?, ?, ?, ?, ?)",
        (old_timestamp, 111, "retention_old", "test", ""),
    )
    pdb._conn.execute(
        "INSERT INTO audit_log (ts, actor_id, action, target, details) "
        "VALUES (?, ?, ?, ?, ?)",
        (recent_timestamp, 111, "retention_recent", "test", ""),
    )
    pdb._conn.execute(
        "INSERT INTO user_messages "
        "(user_id, direction, text, admin_id, created_at) VALUES (?, ?, ?, ?, ?)",
        (777, "in", "رسالة قديمة", None, retention_now.isoformat()),
    )
    old_message_id = pdb._conn.execute(
        "SELECT MAX(id) FROM user_messages"
    ).fetchone()[0]
    pdb._conn.execute(
        "UPDATE user_messages SET created_at = ? WHERE id = ?",
        (retention_now.replace(year=retention_now.year - 1).isoformat(), old_message_id),
    )
    pdb._conn.commit()
retention_counts = asyncio.run(pdb.cleanup_private_data())
assert retention_counts["audit_deleted"] == 1, retention_counts
assert retention_counts["messages_deleted"] == 1, retention_counts
assert any(
    item["action"] == "retention_recent"
    for item in asyncio.run(pdb.recent_audit(100))
)

media_tmp = tempfile.TemporaryDirectory(prefix="mrw-media-retention-test-")
original_media_dir, original_base_dir = config.MEDIA_DIR, config.BASE_DIR
try:
    config.BASE_DIR = Path(media_tmp.name)
    config.MEDIA_DIR = Path(media_tmp.name) / "media"
    config.MEDIA_DIR.mkdir()
    shared_photo = config.MEDIA_DIR / "shared.jpg"
    shared_photo.write_bytes(b"test-image")
    first_shared_report = asyncio.run(pdb.add_report(
        user_id=887, username="shared_user_1", full_name="مستخدم الصورة الأول",
        photo_path="media/shared.jpg", photo_file_id=None,
        post_link="https://example.test/shared-1",
        description="بلاغ يشارك ملف صورة مع بلاغ آخر",
    ))
    second_report = asyncio.run(pdb.add_report(
        user_id=888, username="shared_user", full_name="مستخدم الصورة",
        photo_path="media/shared.jpg", photo_file_id=None,
        post_link="https://example.test/shared",
        description="بلاغ إضافي يشير إلى ملف صورة مشترك",
    ))
    assert shared_photo.exists()
    assert asyncio.run(pdb.delete_report(first_shared_report["id"])) is True
    assert shared_photo.exists(), "shared media must remain while another report references it"
    assert asyncio.run(pdb.delete_report(second_report["id"])) is True
    assert not shared_photo.exists(), "unreferenced report media should be removed"

    evidence_photo = config.MEDIA_DIR / "evidence.jpg"
    evidence_photo.write_bytes(b"evidence")
    asyncio.run(pdb.save_block(
        888, 1, "9999-12-31T00:00:00", "اختبار دليل",
        evidence_photo_path=str(evidence_photo),
    ))
    asyncio.run(pdb.clear_block(888))
    assert not evidence_photo.exists(), "unblocked evidence media should be removed"

    orphan_photo = config.MEDIA_DIR / "orphan.jpg"
    orphan_photo.write_bytes(b"orphan")
    old_mtime = time.time() - 8 * 86400
    os.utime(orphan_photo, (old_mtime, old_mtime))
    assert asyncio.run(pdb.cleanup_orphaned_media()) == 1
    assert not orphan_photo.exists()
finally:
    config.MEDIA_DIR, config.BASE_DIR = original_media_dir, original_base_dir
    media_tmp.cleanup()
print("database: 365/180-day retention and safe orphaned-media cleanup")

# تقدم التعميم محفوظ، وأي مهمة كانت جارية عند إعادة التشغيل تُعلّم كمقاطعة.
jobs_tmp = tempfile.TemporaryDirectory(prefix="mrw-broadcast-selftest-")
jobs_path = Path(jobs_tmp.name) / "jobs.db"
jobs_db = database.Database(jobs_path)
job_id = asyncio.run(jobs_db.create_broadcast_job("تجربة تعميم", 111, 2))
assert job_id is not None
asyncio.run(jobs_db.update_broadcast_job(job_id, 1, 0))
assert asyncio.run(jobs_db.get_broadcast_job(job_id))["sent"] == 1
jobs_db.close()
jobs_db = database.Database(jobs_path)
assert asyncio.run(jobs_db.get_broadcast_job(job_id))["status"] == "interrupted"
job_id = asyncio.run(jobs_db.create_broadcast_job("نشر مكتمل", 111, 1))
assert job_id is not None
assert asyncio.run(jobs_db.finish_broadcast_job(job_id, 1, 0)) == "completed"
assert asyncio.run(jobs_db.get_broadcast_job(job_id))["status"] == "completed"
job_id = asyncio.run(jobs_db.create_broadcast_job("نشر ملغى", 111, 2))
assert job_id is not None
assert asyncio.run(jobs_db.request_broadcast_cancel(job_id))
assert asyncio.run(jobs_db.finish_broadcast_job(job_id, 1, 0)) == "cancelled"
assert asyncio.run(jobs_db.get_broadcast_job(job_id))["status"] == "cancelled"
jobs_db.close()
jobs_tmp.cleanup()

# رفض البلاغ عبر API لا يغيّر الحالة إذا غاب السبب، حتى لو تجاوز تحقق الواجهة.
from fastapi import FastAPI

_bot_main = main_mod

reject_report = asyncio.run(pdb.add_report(
    user_id=909, username="reject_test", full_name="اختبار الرفض",
    photo_path="media/reject.jpg", photo_file_id="reject-file",
    post_link="https://t.me/x/99", description="بلاغ لاختبار إلزام سبب الرفض",
))
api_app = FastAPI()
admin_api.register(api_app)
version_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/version"
)
_old_render_commit = _os.environ.get("RENDER_GIT_COMMIT")
_old_render_environment = _os.environ.get("RENDER_ENVIRONMENT")
_old_render_service_name = _os.environ.get("RENDER_SERVICE_NAME")
try:
    _os.environ["RENDER_GIT_COMMIT"] = "test-commit-sha"
    _os.environ["RENDER_ENVIRONMENT"] = "test-environment"
    version_result = asyncio.run(version_endpoint())
    assert json.loads(version_result.body) == {
        "ok": True,
        "commit": "test-commit-sha",
        "environment": "test-environment",
    }
finally:
    for _key, _old_value in (
        ("RENDER_GIT_COMMIT", _old_render_commit),
        ("RENDER_ENVIRONMENT", _old_render_environment),
        ("RENDER_SERVICE_NAME", _old_render_service_name),
    ):
        if _old_value is None:
            _os.environ.pop(_key, None)
        else:
            _os.environ[_key] = _old_value
status_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/reports/{rid}/status"
)
workflow_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/reports/{rid}/workflow"
)
settings_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/settings"
    and "PUT" in getattr(route, "methods", set())
)
auto_freeze_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/system/auto-freeze"
)
user_limits_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/system/user-limits"
)
followup_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/reports/{rid}/followup"
)
message_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/users/{uid}/message"
)
user_messages_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/users/{uid}/messages"
)
conversations_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/conversations"
)
messages_read_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/users/{uid}/messages/read"
)
retry_user_message_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None)
    == "/api/admin/users/{uid}/messages/{message_id}/retry"
)
user_contact_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/users/{uid}/contact"
)
unblock_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/users/{uid}/unblock"
)
photo_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/photo/{rid}"
)
blocked_user_photo_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/blocked-users/{uid}/photo"
)
blocked_user_report_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/blocked-users/{uid}/report"
)
db_health_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/db-health"
)
system_health_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/system/health"
)
system_integrations_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/system/integrations"
)
osint_sources_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/osint/sources"
    and "GET" in getattr(route, "methods", set())
)
osint_source_add_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/osint/sources"
    and "POST" in getattr(route, "methods", set())
)
osint_source_delete_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/osint/sources/{source_id}"
)
backup_download_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/backup/download"
)
backup_now_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/backup/now"
)
notifications_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/notifications"
)
notification_read_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/notifications/{notification_id}/read"
)
notifications_read_all_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/notifications/read-all"
)
notifications_stats_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/notifications/stats"
)
notifications_cleanup_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/notifications/cleanup"
)
notifications_clear_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/notifications/clear"
)
events_stream_endpoint = next(
    route.endpoint for route in api_app.routes
    if getattr(route, "path", None) == "/api/admin/events/stream"
)
original_bot_app = _bot_main.app
_bot_main.app = SimpleNamespace(bot_data={"db": DB}, bot=None)
db_health_response = db_health_endpoint()
assert db_health_response.status_code == 200
assert json.loads(db_health_response.body) == {"ok": True}
_bot_main.app = original_bot_app
_os.environ["ADMIN_API_TOKEN"] = "test-token-xyz"

class _TursoHealthDb:
    backend = "turso"
    heartbeat = None
    failure_notified = False
    notifications = []

    def health_check(self):
        return True

    async def total(self):
        return 2

    async def get_state(self, key):
        if key == "bot_last_heartbeat":
            return self.heartbeat
        if key == "bot_heartbeat_failure_notified" and self.failure_notified:
            return "1"
        return None

    async def set_state(self, key, value):
        if key == "bot_heartbeat_failure_notified":
            self.failure_notified = value is not None

    async def create_notification(self, notification_type, payload):
        self.notifications.append((notification_type, payload))

class _HealthRequest:
    headers = {"authorization": "Bearer test-token-xyz"}
    method = "GET"
    state = SimpleNamespace()
    url = SimpleNamespace(path="/api/admin/system/health")

_bot_main.app = SimpleNamespace(bot_data={"db": _TursoHealthDb()}, bot=None)
healthy_turso_response = asyncio.run(system_health_endpoint(_HealthRequest()))
healthy_turso = json.loads(healthy_turso_response.body)
assert healthy_turso_response.status_code == 200
assert healthy_turso["database_ready"] is True
assert healthy_turso["backend"] == "turso"
assert healthy_turso["bot_alive"] is False
assert healthy_turso["webhook_ready"] is False
_bot_main.app.bot_data["db"].heartbeat = datetime.now(timezone.utc).isoformat()
fresh_heartbeat_response = asyncio.run(system_health_endpoint(_HealthRequest()))
assert json.loads(fresh_heartbeat_response.body)["bot_alive"] is True
_bot_main.app.bot_data["db"].heartbeat = (
    datetime.now(timezone.utc) - timedelta(seconds=181)
).isoformat()
stale_heartbeat_response = asyncio.run(system_health_endpoint(_HealthRequest()))
assert json.loads(stale_heartbeat_response.body)["bot_alive"] is False
assert _bot_main.app.bot_data["db"].failure_notified is False
assert not _bot_main.app.bot_data["db"].notifications
_bot_main.app.bot_data["db"].heartbeat = None

class _UnavailableTursoDb(_TursoHealthDb):
    def health_check(self):
        raise RuntimeError("simulated database outage")

_bot_main.app = SimpleNamespace(bot_data={"db": _UnavailableTursoDb()}, bot=None)
unavailable_response = asyncio.run(system_health_endpoint(_HealthRequest()))
unavailable_health = json.loads(unavailable_response.body)
assert unavailable_health["database_ready"] is False
assert unavailable_health["backend"] == "turso"

_bot_main.app = SimpleNamespace(bot_data={"db": DB}, bot=None)
healthy_sqlite_response = asyncio.run(system_health_endpoint(_HealthRequest()))
healthy_sqlite = json.loads(healthy_sqlite_response.body)
assert healthy_sqlite["database_ready"] is True
assert healthy_sqlite["backend"] == "sqlite"
assert db_health_endpoint().status_code == 200
print("admin_api: system health checks SELECT 1 and reports backend")

import integrations_monitor

_old_github_activity = integrations_monitor.github_activity
integrations_monitor.github_activity = lambda: {
    "status": "healthy",
    "repository": "ABU-AZZAM-COM/mrw-bot",
    "checked_at": datetime.now(timezone.utc).isoformat(),
    "recent_runs": [{"name": "فحص جودة", "status": "healthy"}],
    "recent_merges": [{"number": 22, "title": "تحسينات لوحة الإدارة"}],
    "error": None,
}
integration_request = _HealthRequest()
integration_request.url = SimpleNamespace(path="/api/admin/system/integrations")
integration_payload = json.loads(
    asyncio.run(system_integrations_endpoint(integration_request)).body
)
assert integration_payload["github"]["status"] == "healthy"
assert any(item["id"] == "database" for item in integration_payload["services"])
responses = iter([
    {
        "workflow_runs": [{
            "name": "فحص الجودة",
            "status": "completed",
            "conclusion": "failure",
            "event": "push",
            "html_url": "https://github.com/example/run/1",
        }]
    },
    [{"number": 23, "title": "تغيير مدمج", "merged_at": "2026-10-25T00:00:00Z"}],
])
activity = integrations_monitor._load_activity(lambda _url: next(responses))
assert activity["status"] == "degraded"
assert activity["recent_merges"][0]["number"] == 23
integrations_monitor.github_activity = _old_github_activity
assert admin_api._sanitize_public_source_url(
    "https://example.org/story?token=private#section"
) == "https://example.org/story"
assert admin_api._sanitize_public_source_url("http://127.0.0.1/private") is None

class _OsintRequest:
    headers = {"x-admin-token": "test-token-xyz"}
    query_params = {}
    state = SimpleNamespace()

    def __init__(self, path, method="GET", body=None):
        self.url = SimpleNamespace(path=path)
        self.method = method
        self.body = body or {}

    async def json(self):
        return self.body


source_request = _OsintRequest("/api/admin/osint/sources", "POST", {
    "report_id": None,
    "url": "https://example.org/story?secret=remove",
    "title": "مصدر عام للتحقق",
    "notes": "تاريخ النشر وسياق الخبر",
    "source_level": "primary",
})
source_payload = json.loads(
    asyncio.run(osint_source_add_endpoint(source_request)).body
)
assert source_payload["item"]["url"] == "https://example.org/story"
source_list_request = _OsintRequest("/api/admin/osint/sources")
source_list = json.loads(
    asyncio.run(osint_sources_endpoint(source_list_request)).body
)
assert len(source_list["items"]) == 1
source_delete_request = _OsintRequest(
    f"/api/admin/osint/sources/{source_payload['item']['id']}",
    "DELETE",
)
source_delete = asyncio.run(osint_source_delete_endpoint(
    source_payload["item"]["id"], source_delete_request
))
assert json.loads(source_delete.body)["deleted"] is True
print("admin_api: read-only integration status and privacy-safe OSINT references")

heartbeat_values = []

class _HeartbeatDb:
    async def set_state(self, key, value):
        heartbeat_values.append((key, value))

asyncio.run(main_mod.write_heartbeat(_HeartbeatDb()))
assert heartbeat_values[0][0] == "bot_last_heartbeat"
heartbeat_written_at = datetime.fromisoformat(heartbeat_values[0][1])
assert abs((datetime.now(timezone.utc) - heartbeat_written_at).total_seconds()) < 5
print("main: heartbeat writes a UTC timestamp to system_state")

class _WebhookBot:
    def __init__(self, fail_set=False):
        self.calls = []
        self.fail_set = fail_set

    async def delete_webhook(self, **kwargs):
        self.calls.append(("delete_webhook", kwargs))

    async def set_webhook(self, **kwargs):
        self.calls.append(("set_webhook", kwargs))
        if self.fail_set:
            raise RuntimeError("simulated Telegram failure")

class _WebhookUpdater:
    def __init__(self):
        self.running = False
        self.calls = []

    async def start_polling(self, **kwargs):
        self.calls.append(("start_polling", kwargs))
        self.running = True

    async def stop(self):
        self.running = False

async def _exercise_webhook_startup(fail_set=False):
    fake_bot = _WebhookBot(fail_set=fail_set)
    fake_updater = _WebhookUpdater() if fail_set else None

    class _WebhookApplication:
        post_init = None
        post_shutdown = None

        def __init__(self):
            self.bot = fake_bot
            self.bot_data = {}
            self.updater = fake_updater
            self.calls = []

        async def initialize(self):
            self.calls.append("initialize")

        async def start(self):
            self.calls.append("start")
            asyncio.get_running_loop().call_soon(
                asyncio.current_task().cancel
            )

        async def stop(self):
            self.calls.append("stop")

        async def shutdown(self):
            self.calls.append("shutdown")

    fake_application = _WebhookApplication()
    try:
        await main_mod.run_webhook(fake_application, "unit-test-secret")
    except asyncio.CancelledError:
        pass
    return fake_application, fake_bot, fake_updater

webhook_run, webhook_bot, _ = asyncio.run(_exercise_webhook_startup())
assert webhook_bot.calls[0] == (
    "delete_webhook", {"drop_pending_updates": False}
)
assert webhook_bot.calls[1][0] == "set_webhook"
assert webhook_bot.calls[1][1]["url"] == main_mod.TELEGRAM_WEBHOOK_URL
assert webhook_bot.calls[1][1]["secret_token"] == "unit-test-secret"
assert webhook_bot.calls[1][1]["drop_pending_updates"] is False
assert webhook_run.bot_data["webhook_ready"] is True

fallback_run, fallback_bot, fallback_updater = asyncio.run(
    _exercise_webhook_startup(fail_set=True)
)
assert fallback_bot.calls[0] == (
    "delete_webhook", {"drop_pending_updates": False}
)
assert fallback_updater.calls == [(
    "start_polling",
    {"allowed_updates": main_mod.Update.ALL_TYPES, "drop_pending_updates": False},
)]
assert fallback_run.bot_data["webhook_ready"] is False
print("main: webhook setup preserves updates and falls back to polling")

import space_app as space_app_mod
from telegram import Update

webhook_app = space_app_mod.build_app()
webhook_endpoint = next(
    route.endpoint for route in webhook_app.routes
    if getattr(route, "path", None) == "/telegram/webhook"
)
old_webhook_secret = _os.environ.get("TELEGRAM_WEBHOOK_SECRET")
try:
    _os.environ["TELEGRAM_WEBHOOK_SECRET"] = "test-webhook-secret"

    class _WebhookRequest:
        def __init__(self, headers, payload):
            self.headers = headers
            self.payload = payload

        async def json(self):
            return self.payload

    async def _test_webhook():
        update_queue = asyncio.Queue()
        running_loop = asyncio.get_running_loop()
        fake_application = SimpleNamespace(
            bot_data={"loop": running_loop},
            bot=object(),
            update_queue=update_queue,
        )
        original_application = _bot_main.app
        _bot_main.app = fake_application
        try:
            denied = await webhook_endpoint(
                _WebhookRequest({}, {"update_id": 10})
            )
            denied_non_ascii = await webhook_endpoint(
                _WebhookRequest(
                    {"X-Telegram-Bot-Api-Secret-Token": "غير-صالح"},
                    {"update_id": 10},
                )
            )
            assert denied.status_code == 403
            assert denied_non_ascii.status_code == 403
            accepted = await webhook_endpoint(
                _WebhookRequest(
                    {"X-Telegram-Bot-Api-Secret-Token": "test-webhook-secret"},
                    {"update_id": 11},
                )
            )
            await asyncio.sleep(0)
            assert accepted.status_code == 200
            queued_update = update_queue.get_nowait()
            assert isinstance(queued_update, Update)
            assert queued_update.update_id == 11
            return (
                denied.status_code,
                denied_non_ascii.status_code,
                accepted.status_code,
            )
        finally:
            _bot_main.app = original_application

    webhook_statuses = asyncio.run(_test_webhook())
    assert webhook_statuses == (403, 403, 200)
finally:
    if old_webhook_secret is None:
        _os.environ.pop("TELEGRAM_WEBHOOK_SECRET", None)
    else:
        _os.environ["TELEGRAM_WEBHOOK_SECRET"] = old_webhook_secret
print("space_app: webhook validates Telegram secret and queues updates")

bio_routes = [
    route for route in webhook_app.routes
    if getattr(route, "path", None) == "/api/set-bio"
]
assert len(bio_routes) == 1 and bio_routes[0].methods == {"POST"}
bio_endpoint = bio_routes[0].endpoint
old_setup_token = _os.environ.get("SETUP_TOKEN")
try:
    _os.environ["SETUP_TOKEN"] = "test-setup-token"

    class _BioRequest:
        headers = {"X-Setup-Token": "wrong-token"}
        query_params = {"key": "test-setup-token"}

    bio_response = asyncio.run(bio_endpoint(_BioRequest()))
    assert bio_response.status_code == 401
finally:
    if old_setup_token is None:
        _os.environ.pop("SETUP_TOKEN", None)
    else:
        _os.environ["SETUP_TOKEN"] = old_setup_token
print("space_app: set-bio is POST-only and rejects query-string credentials")

class _TursoBackupDb:
    backend = "turso"
    path = DB.path

    def export_sqlite_snapshot(self, output_path):
        DB.export_sqlite_snapshot(output_path)

    async def set_state(self, key, value):
        await DB.set_state(key, value)

    async def audit(self, *args):
        return None

backup_request = _HealthRequest()
backup_request.url = SimpleNamespace(path="/api/admin/backup/download")
_bot_main.app = SimpleNamespace(bot_data={"db": _TursoBackupDb()}, bot=None)
backup_response = asyncio.run(backup_download_endpoint(backup_request))
assert backup_response.status_code == 200
assert Path(backup_response.path).is_file()
backup_conn = sqlite3.connect(backup_response.path)
try:
    assert backup_conn.execute("SELECT COUNT(*) FROM reports").fetchone()[0] > 0
finally:
    backup_conn.close()
asyncio.run(backup_response.background())
assert not Path(backup_response.path).exists()

uploaded_snapshot_paths = []

def _fake_upload_backup(snapshot_path):
    snapshot_path = Path(snapshot_path)
    assert snapshot_path.is_file()
    conn = sqlite3.connect(snapshot_path)
    try:
        assert conn.execute("SELECT COUNT(*) FROM reports").fetchone()[0] > 0
    finally:
        conn.close()
    uploaded_snapshot_paths.append(snapshot_path)
    return True

_old_backup_token, _old_backup_repo = config.HF_BACKUP_TOKEN, config.HF_BACKUP_REPO
try:
    config.HF_BACKUP_TOKEN, config.HF_BACKUP_REPO = "test-token", "test/repo"
    with patch("hub_backup.upload_backup", _fake_upload_backup):
        upload_request = _HealthRequest()
        upload_request.method = "POST"
        upload_request.url = SimpleNamespace(path="/api/admin/backup/now")
        upload_response = asyncio.run(backup_now_endpoint(upload_request))
finally:
    config.HF_BACKUP_TOKEN, config.HF_BACKUP_REPO = (
        _old_backup_token,
        _old_backup_repo,
    )
assert json.loads(upload_response.body)["uploaded"] is True
assert len(uploaded_snapshot_paths) == 1
assert not uploaded_snapshot_paths[0].exists()
print("admin_api: Turso backup download and upload create and clean temporary snapshots")

photo_db_tmp = tempfile.TemporaryDirectory(prefix="mrw-photo-selftest-")
photo_db = database.Database(Path(photo_db_tmp.name) / "photo.db")
_bot_main.app = SimpleNamespace(bot_data={"db": photo_db}, bot=None)
photo_path = Path(photo_db_tmp.name) / "legacy-image.png"
photo_path.write_bytes(b"\x89PNG\r\n\x1a\nfake-png-payload")
legacy_image = asyncio.run(photo_db.add_report(
    user_id=778, username="legacy_image", full_name="صورة قديمة",
    photo_path=str(photo_path), photo_file_id=None,
    post_link="https://t.me/x/image", description="بلاغ صورة محفوظة محليًا",
))
assert admin_api._safe_report(legacy_image)["has_photo"] is True
missing_image = {**legacy_image, "photo_path": "missing-image.png"}
assert admin_api._safe_report(missing_image)["has_photo"] is False

class _PhotoRequest:
    headers = {"authorization": "Bearer test-token-xyz"}
    query_params = {"token": "test-token-xyz"}
    method = "GET"
    state = SimpleNamespace()

    def __init__(self, path):
        self.url = SimpleNamespace(path=path)

local_photo_response = asyncio.run(photo_endpoint(
    legacy_image["id"], _PhotoRequest("/api/admin/photo/1")
))
assert local_photo_response.status_code == 200
assert local_photo_response.headers["content-type"] == "image/png"
assert local_photo_response.headers["cache-control"] == "private, no-store"
asyncio.run(photo_db.save_block(
    780,
    0,
    "9999-12-31T00:00:00",
    "كشف آلي",
    evidence_photo_path=str(photo_path),
))
blocked_photo_response = asyncio.run(blocked_user_photo_endpoint(
    780, _PhotoRequest("/api/admin/blocked-users/780/photo")
))
assert blocked_photo_response.status_code == 200
assert blocked_photo_response.headers["content-type"] == "image/png"
blocked_report_response = asyncio.run(blocked_user_report_endpoint(
    780, _PhotoRequest("/api/admin/blocked-users/780/report")
))
assert blocked_report_response.status_code == 200
assert "تقرير حجب مستخدم" in blocked_report_response.body.decode("utf-8")
print("admin_api: local PNG served with correct MIME; legacy reports expose image availability")

remote_image = asyncio.run(photo_db.add_report(
    user_id=779, username="remote_image", full_name="صورة تلغرام",
    photo_path="", photo_file_id="test-file-id",
    post_link="https://t.me/x/remote", description="بلاغ صورة محفوظة في تلغرام",
))
_original_photo_bytes = admin_api._photo_bytes

async def _fake_photo_bytes(_file_id):
    return b"\x89PNG\r\n\x1a\nfake-png-payload"

admin_api._photo_bytes = _fake_photo_bytes
try:
    remote_photo_response = asyncio.run(photo_endpoint(
        remote_image["id"], _PhotoRequest("/api/admin/photo/2")
    ))
    assert remote_photo_response.status_code == 200
    assert remote_photo_response.headers["content-type"] == "image/png"
    asyncio.run(photo_db.save_block(
        781,
        0,
        "9999-12-31T00:00:00",
        "كشف آلي",
        evidence_photo_file_id="test-evidence-file-id",
    ))
    remote_blocked_photo_response = asyncio.run(blocked_user_photo_endpoint(
        781, _PhotoRequest("/api/admin/blocked-users/781/photo")
    ))
    assert remote_blocked_photo_response.status_code == 200
    assert remote_blocked_photo_response.headers["content-type"] == "image/png"
finally:
    admin_api._photo_bytes = _original_photo_bytes
    photo_db.close()
    photo_db_tmp.cleanup()
    _bot_main.app = SimpleNamespace(bot_data={"db": pdb}, bot=None)
print("admin_api: Telegram PNG is served with signature-detected MIME")

class _Request:
    headers = {"x-admin-token": "test-token-xyz"}
    query_params = {}
    method = "GET"

    def __init__(self, body):
        self._body = body
        self.url = SimpleNamespace(path="/api/admin/notifications")
        self.state = SimpleNamespace()

    async def json(self):
        return self._body

    async def is_disconnected(self):
        return False

notifications_api_response = asyncio.run(
    notifications_endpoint(_Request({}), unread=False, limit=20)
)
notifications_api_payload = json.loads(notifications_api_response.body)
assert notifications_api_response.status_code == 200
assert notifications_api_payload["unread_count"] >= 1
assert notifications_api_payload["items"]
async def _first_sse_event():
    response = await events_stream_endpoint(_Request({}))
    iterator = response.body_iterator
    first = await iterator.__anext__()
    pending_event = asyncio.create_task(iterator.__anext__())
    await asyncio.sleep(0.01)
    await pdb.create_notification("report_new", {"seq": "SSE-TEST"})
    second = await pending_event
    await iterator.aclose()
    return first, second

sse_connected, sse_event = asyncio.run(_first_sse_event())
assert sse_connected == ": connected\n\n"
assert "event: notification" in sse_event
assert '"seq": "SSE-TEST"' in sse_event
api_notification_id = notifications_api_payload["items"][0]["id"]
read_api_response = asyncio.run(
    notification_read_endpoint(api_notification_id, _Request({}))
)
assert read_api_response.status_code == 200
read_all_api_response = asyncio.run(notifications_read_all_endpoint(_Request({})))
assert json.loads(read_all_api_response.body)["ok"] is True

cleanup_now = datetime.now(timezone.utc)
cleanup_seed = [
    ("report_new", 1, 8),
    ("report_new", 0, 31),
    ("report_new", 0, 10),
    ("emergency", 1, 60),
    ("heartbeat_fail", 0, 89),
    ("emergency", 0, 91),
    ("heartbeat_fail", 1, 91),
]
seeded_notification_ids = []
with pdb._lock:
    for notification_type, is_read, age_days in cleanup_seed:
        created_at = (cleanup_now - timedelta(days=age_days)).isoformat(
            timespec="seconds"
        )
        cursor = pdb._conn.execute(
            "INSERT INTO notifications(type, payload, created_at, is_read) "
            "VALUES (?, '{}', ?, ?)",
            (notification_type, created_at, is_read),
        )
        seeded_notification_ids.append(int(cursor.lastrowid))
    pdb._conn.commit()
stats_response = asyncio.run(notifications_stats_endpoint(_Request({})))
stats_payload = json.loads(stats_response.body)
assert stats_response.status_code == 200
assert stats_payload["total"] >= len(seeded_notification_ids)
assert stats_payload["unread"] >= 1
assert stats_payload["oldest_created_at"] is not None
assert stats_payload["oldest_read_created_at"] is not None
cleanup_request = _Request({})
cleanup_request.method = "POST"
cleanup_request.url = SimpleNamespace(path="/api/admin/notifications/cleanup")
cleanup_response = asyncio.run(notifications_cleanup_endpoint(cleanup_request))
cleanup_payload = json.loads(cleanup_response.body)
assert cleanup_response.status_code == 200
assert cleanup_payload == {"ok": True, "deleted": 4}
with pdb._lock:
    retained_ids = {
        int(row["id"])
        for row in pdb._conn.execute(
            "SELECT id FROM notifications WHERE id IN ("
            + ",".join("?" for _ in seeded_notification_ids)
            + ")",
            seeded_notification_ids,
        ).fetchall()
    }
assert retained_ids == set(seeded_notification_ids) - {
    seeded_notification_ids[0],
    seeded_notification_ids[1],
    seeded_notification_ids[5],
    seeded_notification_ids[6],
}
old_admin_tokens_json = _os.environ.get("ADMIN_API_TOKENS_JSON")
try:
    _os.environ["ADMIN_API_TOKENS_JSON"] = json.dumps([
        {"token": "viewer-test-token", "role": "viewer", "admin_id": 22}
    ])
    viewer_cleanup_request = _Request({})
    viewer_cleanup_request.headers = {
        "authorization": "Bearer viewer-test-token"
    }
    viewer_cleanup_request.method = "POST"
    viewer_cleanup_request.url = SimpleNamespace(
        path="/api/admin/notifications/cleanup"
    )
    viewer_cleanup = asyncio.run(
        notifications_cleanup_endpoint(viewer_cleanup_request)
    )
    assert viewer_cleanup.status_code == 403
    viewer_clear_request = _Request({})
    viewer_clear_request.headers = viewer_cleanup_request.headers
    viewer_clear_request.method = "POST"
    viewer_clear_request.url = SimpleNamespace(
        path="/api/admin/notifications/clear"
    )
    viewer_clear = asyncio.run(
        notifications_clear_endpoint(viewer_clear_request)
    )
    assert viewer_clear.status_code == 403, (
        viewer_clear.status_code,
        viewer_clear.body,
    )
finally:
    if old_admin_tokens_json is None:
        _os.environ.pop("ADMIN_API_TOKENS_JSON", None)
    else:
        _os.environ["ADMIN_API_TOKENS_JSON"] = old_admin_tokens_json

clear_db = database.Database(Path(pbase_tmp.name) / "clear-notifications.db")
for notification_type in ("report_new", "emergency", "heartbeat_fail"):
    asyncio.run(clear_db.create_notification(notification_type, {}))
clear_now = datetime.now(timezone.utc)
with clear_db._lock:
    clear_rows = clear_db._conn.execute(
        "SELECT id, type FROM notifications ORDER BY id"
    ).fetchall()
    for row in clear_rows:
        age_days = 30 if row["type"] == "emergency" else (
            91 if row["type"] == "heartbeat_fail" else 0
        )
        created_at = (clear_now - timedelta(days=age_days)).isoformat(
            timespec="seconds"
        )
        clear_db._conn.execute(
            "UPDATE notifications SET created_at = ? WHERE id = ?",
            (created_at, row["id"]),
        )
    clear_db._conn.commit()
old_clear_app = _bot_main.app
_bot_main.app = SimpleNamespace(bot_data={"db": clear_db}, bot=None)
clear_request = _Request({})
clear_request.method = "POST"
clear_request.url = SimpleNamespace(path="/api/admin/notifications/clear")
clear_response = asyncio.run(notifications_clear_endpoint(clear_request))
assert clear_response.status_code == 200
assert json.loads(clear_response.body) == {"ok": True, "deleted": 2}
clear_stats = asyncio.run(clear_db.notification_stats())
assert clear_stats["total"] == 1
with clear_db._lock:
    preserved_types = {
        row["type"]
        for row in clear_db._conn.execute(
            "SELECT type FROM notifications"
        ).fetchall()
    }
assert preserved_types == {"emergency"}
_bot_main.app = old_clear_app
clear_db.close()
print("admin_api: notification cleanup, clear, and 90-day emergency retention")

response = asyncio.run(status_endpoint(
    reject_report["id"], _Request({"status": "rejected"})
))
assert response.status_code == 400
assert json.loads(response.body)["error"] == "reason_required"
assert asyncio.run(pdb.get_by_id(reject_report["id"]))["status"] == database.STATUS_NEW

due_at = (datetime.now(timezone.utc) + timedelta(days=1)).isoformat()
workflow_response = asyncio.run(workflow_endpoint(
    reject_report["id"],
    _Request({
        "priority": "high",
        "assigned_to": 111111111,
        "due_at": due_at,
    }),
))
assert workflow_response.status_code == 200
workflow_saved = asyncio.run(pdb.get_by_id(reject_report["id"]))
assert workflow_saved["priority"] == "high"
assert workflow_saved["assigned_to"] == 111111111

old_runtime_settings = {
    name: getattr(config, name) for name in config.RUNTIME_SETTING_LIMITS
}
settings_response = asyncio.run(settings_endpoint(
    _Request({
        "max_description_length": 1800,
        "max_reports_per_hour": 7,
        "max_reports_per_hour_global": 120,
        "max_reports_per_day_global": 600,
        "freeze_after_strikes": 2,
        "freeze_minutes_first": 15,
        "freeze_minutes_second": 120,
        "freeze_minutes_repeat": 720,
        "block_messages": {
            "frozen": "اختبار حظر: {minutes} دقيقة",
            "manual_block": "تم تقييد الحساب: {kind} — {reason}",
        },
    })
))
assert settings_response.status_code == 200
assert config.MAX_DESC_LEN == 1800
assert config.MAX_REPORTS_PER_HOUR == 7
assert config.MAX_REPORTS_PER_HOUR_GLOBAL == 120
assert config.MAX_REPORTS_PER_DAY_GLOBAL == 600
assert config.FREEZE_AFTER_STRIKES == 2
assert config.FREEZE_MINUTES_FIRST == 15
assert config.FREEZE_MINUTES_SECOND == 120
assert config.FREEZE_MINUTES_REPEAT == 720
saved_messages = json.loads(settings_response.body)["items"]["block_messages"]
assert saved_messages["frozen"] == "اختبار حظر: {minutes} دقيقة"
assert saved_messages["manual_block"] == "تم تقييد الحساب: {kind} — {reason}"
bad_message_response = asyncio.run(settings_endpoint(
    _Request({"block_messages": {"frozen": "تجاوز غير مدعوم {unknown}"}})
))
assert bad_message_response.status_code == 400
bad_order_response = asyncio.run(settings_endpoint(
    _Request({
        "freeze_minutes_first": 500,
        "freeze_minutes_second": 200,
        "freeze_minutes_repeat": 700,
    })
))
assert bad_order_response.status_code == 400
config.apply_runtime_settings(old_runtime_settings)
rendered_message = asyncio.run(
    handlers.guards.render_block_message(pdb, "frozen", {"minutes": "<12>"})
)
assert "&lt;12&gt;" in rendered_message and "<12>" not in rendered_message

invalid_pause = asyncio.run(auto_freeze_endpoint(
    _Request({"on": True, "duration_minutes": 4, "reason": "اختبار"})
))
assert invalid_pause.status_code == 400
assert asyncio.run(pdb.get_state("auto_freeze_paused_until")) is None
pause_response = asyncio.run(auto_freeze_endpoint(
    _Request({"on": True, "duration_minutes": 30, "reason": "اختبار لوحة"})
))
assert pause_response.status_code == 200
assert json.loads(pause_response.body)["paused"] is True
resume_response = asyncio.run(auto_freeze_endpoint(
    _Request({"on": False, "reason": "انتهاء الاختبار"})
))
assert resume_response.status_code == 200
assert asyncio.run(pdb.get_state("auto_freeze_paused_until")) is None

limits_pause_response = asyncio.run(user_limits_endpoint(
    _Request({"on": True, "duration_minutes": 60, "reason": "اختبار لوحة"})
))
assert limits_pause_response.status_code == 200
assert asyncio.run(pdb.get_state("user_limits_paused_until"))
limits_resume_response = asyncio.run(user_limits_endpoint(
    _Request({"on": False, "reason": "انتهاء الاختبار"})
))
assert limits_resume_response.status_code == 200
assert asyncio.run(pdb.get_state("user_limits_paused_until")) is None

send_calls = []
send_results = [{"sent": True, "message_id": 7001}]

async def _fake_send_direct_message(user_id, text):
    send_calls.append((user_id, text))
    return True

async def _fake_send_direct_message_result(user_id, text):
    send_calls.append((user_id, text))
    return send_results.pop(0)

original_sender = admin_api._send_direct_message
original_detailed_sender = admin_api._send_direct_message_result
admin_api._send_direct_message = _fake_send_direct_message
admin_api._send_direct_message_result = _fake_send_direct_message_result
followup_response = asyncio.run(followup_endpoint(
    reject_report["id"], _Request({"note": "أرسل تفاصيل إضافية"})
))
assert followup_response.status_code == 200
followup_saved = asyncio.run(pdb.get_by_id(reject_report["id"]))
assert followup_saved["followup_requested_at"]
assert len(send_calls) == 1 and "/reply" in send_calls[0][1]
message_response = asyncio.run(message_endpoint(
    777, _Request({"text": "رسالة اختبار"})
))
assert message_response.status_code == 200
assert json.loads(message_response.body)["sent"] is True
assert send_calls[0][0] == 909 and "/reply" in send_calls[0][1]
assert send_calls[1] == (777, "رسالة اختبار")
messages_response = asyncio.run(user_messages_endpoint(777, _Request({})))
messages_payload = json.loads(messages_response.body)
assert messages_response.status_code == 200
assert messages_payload["items"][-1]["text"] == "رسالة اختبار"
assert messages_payload["items"][-1]["delivery_status"] == "delivered"
conversation_payload = json.loads(
    asyncio.run(conversations_endpoint(_Request({}))).body
)
conversation = next(item for item in conversation_payload["items"] if item["user_id"] == 777)
assert conversation["last_message"] == "رسالة اختبار"
assert conversation["unread_count"] >= 1
asyncio.run(messages_read_endpoint(777, _Request({})))
conversation_payload = json.loads(
    asyncio.run(conversations_endpoint(_Request({}))).body
)
conversation = next(item for item in conversation_payload["items"] if item["user_id"] == 777)
assert conversation["unread_count"] == 0
asyncio.run(pdb.add_user_message(777, "in", "رسالة واردة"))
conversation_payload = json.loads(
    asyncio.run(conversations_endpoint(_Request({}))).body
)
conversation = next(item for item in conversation_payload["items"] if item["user_id"] == 777)
assert conversation["unread_count"] == 1
inbound_messages = json.loads(
    asyncio.run(user_messages_endpoint(777, _Request({}))).body
)["items"]
inbound_message_id = max(
    item["id"] for item in inbound_messages if item["direction"] == "in"
)
stale_read_response = asyncio.run(
    messages_read_endpoint(777, _Request({"last_message_id": 0}))
)
assert stale_read_response.status_code == 200
conversation_payload = json.loads(
    asyncio.run(conversations_endpoint(_Request({}))).body
)
conversation = next(item for item in conversation_payload["items"] if item["user_id"] == 777)
assert conversation["unread_count"] == 1
invalid_read_response = asyncio.run(
    messages_read_endpoint(777, _Request({"last_message_id": "invalid"}))
)
assert invalid_read_response.status_code == 400
read_payload = json.loads(
    asyncio.run(
        messages_read_endpoint(
            777, _Request({"last_message_id": inbound_message_id})
        )
    ).body
)
assert read_payload["last_read_message_id"] == inbound_message_id
conversation_payload = json.loads(
    asyncio.run(conversations_endpoint(_Request({}))).body
)
conversation = next(item for item in conversation_payload["items"] if item["user_id"] == 777)
assert conversation["unread_count"] == 0
send_results.append({"sent": False, "error_code": "user_blocked_bot"})
failed_message_response = asyncio.run(message_endpoint(
    777, _Request({"text": "محاولة فاشلة للاختبار"})
))
assert failed_message_response.status_code == 502
failed_payload = json.loads(failed_message_response.body)
assert failed_payload["error"] == "user_blocked_bot"
failed_message_request = _Request({})
failed_message_request.query_params = {
    "after_id": str(failed_payload["message_id"] - 1)
}
failed_message = asyncio.run(user_messages_endpoint(777, failed_message_request))
failed_item = json.loads(failed_message.body)["items"][0]
assert failed_item["delivery_status"] == "failed"
assert failed_item["error_code"] == "user_blocked_bot"
send_results.append({"sent": True, "message_id": 7002})
retry_response = asyncio.run(retry_user_message_endpoint(
    777, failed_payload["message_id"], _Request({})
))
assert retry_response.status_code == 200
assert json.loads(retry_response.body)["sent"] is True
retried_message_request = _Request({})
retried_message_request.query_params = {
    "after_id": str(failed_payload["message_id"] - 1)
}
retried_items = asyncio.run(user_messages_endpoint(777, retried_message_request))
assert json.loads(retried_items.body)["items"][0]["delivery_status"] == "delivered"
assert any(
    row["action"] == "user_message" and row["target"] == "777"
    for row in asyncio.run(pdb.recent_audit(20))
)
contact_request = _Request({"enabled": False})
contact_request.method = "PUT"
contact_request.url = SimpleNamespace(path="/api/admin/users/777/contact")
contact_response = asyncio.run(user_contact_endpoint(777, contact_request))
contact_payload = json.loads(contact_response.body)
assert contact_response.status_code == 200
assert contact_payload == {
    "ok": True, "enabled": False, "changed": True, "notified": True
}
assert send_calls[-1][0] == 777 and "إيقاف" in send_calls[-1][1]
disabled_message = _Request({"text": "رسالة يجب منعها"})
disabled_message.url = SimpleNamespace(path="/api/admin/users/777/message")
assert asyncio.run(message_endpoint(777, disabled_message)).status_code == 403
followup_contact_report = asyncio.run(pdb.add_report(
    user_id=909, username="followup_user", full_name="مستخدم المتابعة",
    photo_path="", photo_file_id=None, post_link="https://example.test/followup",
    description="بلاغ مفتوح لاختبار إعداد التواصل المباشر",
))
asyncio.run(pdb.set_user_contact_enabled(909, False))
send_count_before_blocked_followup = len(send_calls)
blocked_followup = asyncio.run(followup_endpoint(
    followup_contact_report["id"], _Request({"note": "متابعة يجب منعها"})
))
assert blocked_followup.status_code == 403
assert len(send_calls) == send_count_before_blocked_followup
contact_request = _Request({"enabled": True})
contact_request.method = "PUT"
contact_request.url = SimpleNamespace(path="/api/admin/users/777/contact")
assert json.loads(asyncio.run(user_contact_endpoint(777, contact_request)).body)[
    "enabled"
] is True
admin_api._send_direct_message = original_sender
admin_api._send_direct_message_result = original_detailed_sender
notices = []

async def _fake_unblock_notice(user_id):
    notices.append(user_id)
    return True

original_unblock_notice = admin_api._send_unblock_notice
admin_api._send_unblock_notice = _fake_unblock_notice
asyncio.run(pdb.save_block(777, 3, "9999-12-31T00:00:00", "اختبار رفع الحجب"))
unblock_response = asyncio.run(unblock_endpoint(777, _Request({})))
unblock_payload = json.loads(unblock_response.body)
assert unblock_response.status_code == 200
assert unblock_payload["notified"] is True and notices == [777]
assert asyncio.run(pdb.get_block(777))["strikes"] == 0
assert "اختبار رفع الحجب" in asyncio.run(pdb.recent_audit(1))[0]["details"]
admin_api._send_unblock_notice = original_unblock_notice
_os.environ.pop("ADMIN_API_TOKEN", None)
_bot_main.app = original_bot_app
print("admin_api: رفض الحالة بلا سبب مرفوض بلا تعديل على البلاغ")
print("admin_api: الرسالة المباشرة تُرسل عبر طبقة البوت وتُسجل")

# الحذف النهائي يعيد True ثم يختفي البلاغ
assert asyncio.run(pdb.delete_report(pid["id"])) is True
assert asyncio.run(pdb.get_by_id(pid["id"])) is None

# --- التغطية الشاملة: تصنيفات/قوالب/إحصاءات/تجميد (منطق القاعدة) ---
# التصنيفات: الافتراضي من الإعداد ثم تعديل فعل ثم رجوع
import config as _cfg_mod

_default = asyncio.run(pdb.effective_vtypes())
assert list(_default) == list(_cfg_mod.VIOLATION_TYPES)
await_none = asyncio.run(pdb.set_vtypes_override(["نوع أ", "نوع ب", "نوع ج"]))
assert asyncio.run(pdb.effective_vtypes()) == ["نوع أ", "نوع ب", "نوع ج"]
await_none = asyncio.run(pdb.set_vtypes_override([]))
assert asyncio.run(pdb.effective_vtypes()) == list(_cfg_mod.VIOLATION_TYPES)

# القوالب الجاهزة: حفظ/قراءة/فراغ
await_none = asyncio.run(pdb.set_templates(["تم قبول بلاغك.", "بلاغك مرفوض لأسباب الشروط."]))
tpl = asyncio.run(pdb.get_templates())
assert tpl == ["تم قبول بلاغك.", "بلاغك مرفوض لأسباب الشروط."]

# إحصاءات: بلاغان فاعلان ثم يوميات/أنواع/أعلى المتبلّغين/شرائح المستخدمين
await_none = asyncio.run(pdb.add_report(
    user_id=1888888001, username="st8", full_name="مستخدم إحصاء",
    photo_path="m/x.jpg", photo_file_id="f8",
    post_link="https://t.me/x/22", description="وصف إحصائي",
    violation_type="مخالفة زي",
))
daily = asyncio.run(pdb.count_daily("2000-01-01"))
assert daily and daily[-1]["n"] >= 1 and len(daily[-1]["day"]) == 10
by_type = asyncio.run(pdb.count_by_type())
assert by_type and by_type[0]["vtype"] == "مخالفة زي"
tops = asyncio.run(pdb.top_users(7))
stats_user = next(row for row in tops if row["user_id"] == 1888888001)
assert stats_user["n"] == 1

async def _concurrent_analytics_reads():
    async def read_snapshot():
        daily_rows, type_rows, user_rows = await asyncio.gather(
            pdb.count_daily("2000-01-01"),
            pdb.count_by_type(),
            pdb.top_users(7),
        )
        assert isinstance(daily_rows, list)
        assert isinstance(type_rows, list)
        assert isinstance(user_rows, list)

    async def add_analytics_report(index):
        await pdb.add_report(
            user_id=9000 + index,
            username=f"parallel_{index}",
            full_name=f"مستخدم متزامن {index}",
            photo_path="",
            photo_file_id=None,
            post_link=f"https://example.com/parallel/{index}",
            description=f"اختبار التزامن {index}",
            violation_type="اختبار التزامن",
        )

    await asyncio.gather(
        *(read_snapshot() for _ in range(12)),
        *(add_analytics_report(index) for index in range(8)),
    )

asyncio.run(_concurrent_analytics_reads())
assert asyncio.run(pdb.count_by_type())[0]["n"] == 8
print("analytics queries remain safe during concurrent Turso-style access")

ids, overflow = asyncio.run(pdb.all_user_ids(cap=300))
assert 888 in ids and overflow is False
ids, overflow = asyncio.run(pdb.all_user_ids(cap=1))
assert len(ids) <= 1

# مسار القاعدة للتنزيل
assert pdb.path.is_file()
pdb.close()
pbase_tmp.cleanup()
print("admin_api: أدوات اللوحة على القاعدة (بحث/تصنيف/ملاحظة/حجب/حذف) OK")
print("admin_api: تصنيفات ديناميكية + قوالب + إحصاءات + معرفات التعميم OK")

# خادم الصحة: ليس بيئة سحابية محليًا
assert health_server.is_space_environment() is False
print("health_server OK")

print("\n=== كل الاختبارات نجحت ✅ ===")
DB.close()
_db_test_dir.cleanup()
