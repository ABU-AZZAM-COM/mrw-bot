"""
تصدير البلاغات إلى ملف Excel (pandas + openpyxl) و ملف نصي منظم.

- ملفات التصدير تُحفظ في مجلد exports/ ثم يرسلها البوت كمستند.
- حماية من حقن الصيغ (Formula Injection): أي قيمة تبدأ بـ = + - @ أو Tab
  يُضاف قبلها فاصلة عليا حتى تظل نصًا داخل Excel.
"""
from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

import config
import utils
from database import STATUS_NEW, STATUS_REVIEW, STATUS_APPROVED, STATUS_REJECTED

# ترتيب أعمدة قاعدة البيانات الداخلية -> الاسم العربي المعروض في Excel
EXCEL_COLS: list[tuple[str, str]] = [
    ("seq",         "رقم البلاغ"),
    ("status",      "الحالة"),
    ("full_name",   "اسم المرسل"),
    ("username",    "معرف المرسل"),
    ("user_id",     "الرقم التعريفي"),
    ("created_at",  "وقت الإرسال"),
    ("post_link",   "رابط المنشور"),
    ("description", "وصف المخالفة"),
    ("photo_path",  "الصورة"),
    ("reviewed_by", "راجعه (ID)"),
    ("reviewed_at", "وقت المراجعة"),
    ("admin_note",  "ملاحظات المشرف"),
]

_WIDTHS = {
    "seq": 16, "status": 15, "full_name": 22, "username": 18, "user_id": 14,
    "created_at": 20, "post_link": 45, "description": 70, "photo_path": 35,
    "reviewed_by": 14, "reviewed_at": 20, "admin_note": 30,
}

_FORMULA_PREFIXES = ("=", "+", "@", "-", "\t", "\r")
_ILLEGAL_EXCEL_CHARACTERS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _formula_safe(value):
    """تحييد أي قيمة قد تُفسّر كصيغة داخل Excel."""
    if isinstance(value, str):
        value = _ILLEGAL_EXCEL_CHARACTERS.sub("", value)
        if value[:1] in _FORMULA_PREFIXES:
            return "'" + value
    return value


def _cell_value(row: dict, key: str):
    """تحويل قيمة الحقل إلى النسق المناسب للعرض."""
    v = row.get(key)
    if v is None:
        return "-"
    if key in ("created_at", "reviewed_at"):
        return utils.fmt_local(v)
    if key == "username":
        return f"@{v}" if v else "بدون معرف"
    if key == "photo_path":
        try:
            return Path(str(v)).name
        except Exception:
            return str(v)
    return v if isinstance(v, (int, float)) else str(v)


def _stamp() -> str:
    return dt.datetime.now().strftime("%Y%m%d-%H%M%S")


# ==================================================================
# تصدير Excel
# ==================================================================
def export_excel(rows: list[dict], out_dir: Path | None = None) -> Path:
    """إنشاء ملف Excel منسق (اتجاه يمين-يسار، رأس ملوّن، روابط قابلة للنقر)."""
    import pandas as pd
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    out_dir = Path(out_dir or config.EXPORTS_DIR)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"reports_export_{_stamp()}.xlsx"

    data = [
        {ar: _formula_safe(_cell_value(r, key)) for key, ar in EXCEL_COLS}
        for r in rows
    ]
    df = pd.DataFrame(data, columns=[ar for _, ar in EXCEL_COLS])

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="البلاغات")
        ws = writer.sheets["البلاغات"]

        ws.sheet_view.rightToLeft = True
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        # تنسيق صف العناوين
        fill = PatternFill("solid", fgColor="1F4E78")
        head_font = Font(bold=True, color="FFFFFF")
        head_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for cell in ws[1]:
            cell.fill = fill
            cell.font = head_font
            cell.alignment = head_align

        # ضبط عرض الأعمدة
        for idx, (key, _) in enumerate(EXCEL_COLS, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = _WIDTHS.get(key, 20)

        # وصف المخالفة: التفاف نص + محاذاة علوية
        desc_idx = [k for k, _ in EXCEL_COLS].index("description") + 1
        for row_i in range(2, ws.max_row + 1):
            ws.cell(row=row_i, column=desc_idx).alignment = Alignment(
                vertical="top", wrap_text=True
            )

        # جعل روابط المنشور قابلة للنقر
        link_idx = [k for k, _ in EXCEL_COLS].index("post_link") + 1
        link_font = Font(color="0563C1", underline="single")
        for row_i in range(2, ws.max_row + 1):
            cell = ws.cell(row=row_i, column=link_idx)
            val = cell.value
            if isinstance(val, str) and val.startswith(("http://", "https://")):
                cell.hyperlink = val
                cell.font = link_font

        # ورقة "ملخص" بعدد البلاغات لكل حالة
        counts = {s: 0 for s in (STATUS_NEW, STATUS_REVIEW, STATUS_APPROVED, STATUS_REJECTED)}
        for r in rows:
            counts[r["status"]] = counts.get(r["status"], 0) + 1

        summary = writer.book.create_sheet("ملخص")
        summary.sheet_view.rightToLeft = True
        summary.append(["البند", "العدد"])
        summary.append(["الإجمالي", len(rows)])
        summary.append([STATUS_NEW, counts.get(STATUS_NEW, 0)])
        summary.append([STATUS_REVIEW, counts.get(STATUS_REVIEW, 0)])
        summary.append([STATUS_APPROVED, counts.get(STATUS_APPROVED, 0)])
        summary.append([STATUS_REJECTED, counts.get(STATUS_REJECTED, 0)])
        for c in summary["A1:B1"][0]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = fill

    return path


# ==================================================================
# تصدير نصي منظم
# ==================================================================
def export_text(rows: list[dict], out_dir: Path | None = None) -> Path:
    """إنشاء ملف نصي منظم (UTF-8 مع BOM ليتفتح سليمًا في ويندوز)."""
    out_dir = Path(out_dir or config.EXPORTS_DIR)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"reports_export_{_stamp()}.txt"

    lines: list[str] = []
    bar = "=" * 72
    sep = "-" * 72

    counts = {STATUS_NEW: 0, STATUS_REVIEW: 0, STATUS_APPROVED: 0, STATUS_REJECTED: 0}
    for r in rows:
        counts[r["status"]] = counts.get(r["status"], 0) + 1

    lines.append(bar)
    lines.append("     " + "سجل بلاغات المخالفات العسكرية")
    lines.append(f"     تاريخ إنشاء الملف: {utils.fmt_local(utils.utc_now_iso())}")
    lines.append(
        f"     عدد البلاغات: {len(rows)} "
        f"(جديد: {counts['جديد']} | قيد المراجعة: {counts['قيد المراجعة']} | "
        f"معتمد: {counts['معتمد']} | مرفوض: {counts['مرفوض']})"
    )
    lines.append(bar)
    lines.append("")

    for i, r in enumerate(rows, start=1):
        username = f"@{r['username']}" if r["username"] else "بدون معرف"
        reviewed_at = utils.fmt_local(r.get("reviewed_at"))
        lines.append(f"[{i}] بلاغ رقم: {r['seq']}")
        lines.append(f"الحالة          : {r['status']}")
        name = r["full_name"] or "-"
        lines.append(f"المرسل          : {name}")
        lines.append(f"المعرف / المُعرّف: {username} (ID: {r['user_id']})")
        lines.append(f"وقت الإرسال     : {utils.fmt_local(r['created_at'])}")
        lines.append(f"رابط المنشور    : {r['post_link']}")
        lines.append("وصف المخالفة    :")
        for ln in str(r["description"]).splitlines() or ["-"]:
            lines.append(f"    {ln}")
        if r.get("reviewed_by"):
            lines.append(
                f"المراجعة        : بواسطة المشرف {r['reviewed_by']} في {reviewed_at}"
            )
        else:
            lines.append("المراجعة        : لم تُراجع بعد")
        if r.get("photo_path"):
            lines.append(f"الصورة          : {Path(str(r['photo_path'])).name}")
        lines.append(sep)
        lines.append("")

    body = "\n".join(lines)
    with open(path, "w", encoding="utf-8-sig", newline="\n") as f:
        f.write(body)
    return path
